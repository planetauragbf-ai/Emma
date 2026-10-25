# -*- coding: utf-8 -*-
"""Assemble la soutenance en PPTX 16:9 : une image par slide + notes orales minutées."""
import json, glob, os
from pptx import Presentation
from pptx.util import Emu
HERE = os.path.dirname(os.path.abspath(__file__))
notes = json.load(open(os.path.join(HERE, "notes.json")))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
for i, png in enumerate(sorted(glob.glob(os.path.join(HERE, "png_ExportPilot360_Soutenance", "*.png")))):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    s.notes_slide.notes_text_frame.text = notes[i] if i < len(notes) else ""
out = os.path.join(HERE, "..", "decks", "ExportPilot360_Soutenance.pptx")
prs.save(out)
print("pptx", len(prs.slides.__iter__.__self__._sldIdLst), out)

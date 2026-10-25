# -*- coding: utf-8 -*-
"""Construit le modèle pour un scénario, le recalcule avec LibreOffice et exporte les résultats en JSON."""
import json, subprocess, sys, os, shutil, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))

def run(scen, out_xlsx, overrides=None):
    tmp = tempfile.mkdtemp()
    raw = os.path.join(tmp, "raw.xlsx")
    subprocess.run([sys.executable, os.path.join(HERE, "build_model.py"), str(scen), raw], check=True, capture_output=True)
    mp = json.load(open(raw.replace(".xlsx", "_map.json")))
    if overrides:
        wb = load_workbook(raw)
        for name, val in overrides.items():
            sh, cell = mp["H"][name].split("!")
            wb[sh][cell.replace("$", "")] = val
        wb.save(raw)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "out"), raw],
                   check=True, capture_output=True)
    rec = os.path.join(tmp, "out", "raw.xlsx")
    shutil.copy(rec, out_xlsx)
    wb = load_workbook(rec, data_only=True)
    sy, ms = wb["Synthèse"], wb["Mensuel"]
    res = {"synthese": {}, "mensuel": {}}
    for name, r in mp["S"].items():
        res["synthese"][name] = {c: sy[f"{c}{r}"].value for c in "BCDE"}
    for name, r in mp["R"].items():
        res["mensuel"][name] = [ms.cell(row=r, column=3 + i).value for i in range(mp["N"])]
    return res

if __name__ == "__main__":
    allres = {}
    for scen, label in [(1, "pessimiste"), (2, "central"), (3, "optimiste")]:
        out = os.path.join(HERE, "Modele_Financier_ExportPilot360.xlsx" if scen == 2 else f"_scenario_{label}.xlsx")
        allres[label] = run(scen, out)
    # besoin de financement brut : central sans aucune ressource externe
    allres["central_sans_financement"] = run(2, os.path.join(HERE, "_central_sans_financement.xlsx"),
        {"apport": 0, "honneur": 0, "bft1": 0, "bft2": 0, "bank": 0, "ba": 0})
    # pessimiste avec plan de gel : pas de développeur, ni de 2e BizDev, ni de customer success sur la période
    allres["pessimiste_plan_gel"] = run(1, os.path.join(HERE, "_scenario_pessimiste_plan_gel.xlsx"),
        {"start4": 99, "start5": 99, "start6": 99})
    json.dump(allres, open(os.path.join(HERE, "resultats.json"), "w"), ensure_ascii=False, indent=1, default=str)
    print("ok")

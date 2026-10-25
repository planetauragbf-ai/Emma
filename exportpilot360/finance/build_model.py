# -*- coding: utf-8 -*-
"""ExportPilot 360 — modèle financier 3 ans (ante-création nov.-déc. 2026, exercices 2027-2029).

Génère Modele_Financier_ExportPilot360.xlsx, 100 % formules :
  Hypothèses (bleu = modifiable)  →  Mensuel (moteur, 38 mois)  →  Synthèse (CR, financement, KPI, seuil).
Le sélecteur de scénario (Hypothèses!B4 : 1 pessimiste, 2 central, 3 optimiste) pilote
les recrutements de clients et le churn. Recalcul via LibreOffice (voir recalc.py).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

SCENARIO = int(sys.argv[1]) if len(sys.argv) > 1 else 2
OUT = sys.argv[2] if len(sys.argv) > 2 else "Modele_Financier_ExportPilot360.xlsx"

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True, color="3A1430")
SEC = Font(name=F, size=10, bold=True, color="FFFFFF")
NOTE = Font(name=F, size=9, italic=True, color="666666")
SEC_FILL = PatternFill("solid", fgColor="3A1430")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
TOT_FILL = PatternFill("solid", fgColor="EFE7EC")
EUR = '#,##0;(#,##0);"-"'
PCT = '0.0%'
NB1 = '#,##0.0'
NB = '#,##0'

wb = Workbook()

# =========================================================== HYPOTHÈSES
hy = wb.active
hy.title = "Hypothèses"
hy.column_dimensions["A"].width = 58
for c in "BCDEFGHIJKLMNO":
    hy.column_dimensions[c].width = 12
H = {}   # nom -> référence absolue
row = [1]

def title(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = TITLE

def section(ws, text, ncols=6):
    r = row[0] + 1
    for c in range(1, ncols + 1):
        ws.cell(row=r, column=c).fill = SEC_FILL
    ws.cell(row=r, column=1, value=text).font = SEC
    row[0] = r + 1

def hyp(name, label, value, fmt=EUR, note=None, key=False):
    r = row[0]
    hy.cell(row=r, column=1, value=label).font = BLACK
    c = hy.cell(row=r, column=2, value=value)
    c.font = BLUE if not (isinstance(value, str) and value.startswith("=")) else BLACK
    c.number_format = fmt
    if key:
        c.fill = KEY_FILL
    if note:
        hy.cell(row=r, column=3, value=note).font = NOTE
    H[name] = f"Hypothèses!$B${r}"
    row[0] = r + 1
    return r

def hyp_years(name, label, values, fmt=EUR, note=None):
    """Valeurs 2027, 2028, 2029 en colonnes C, D, E."""
    r = row[0]
    hy.cell(row=r, column=1, value=label).font = BLACK
    for i, v in enumerate(values):
        c = hy.cell(row=r, column=3 + i, value=v)
        c.font = BLUE
        c.number_format = fmt
    if note:
        hy.cell(row=r, column=7, value=note).font = NOTE
    H[name] = r
    row[0] = r + 1
    return r

title(hy, 1, "EXPORTPILOT 360 — HYPOTHÈSES DU BUSINESS PLAN")
hy.cell(row=2, column=1, value="Bleu = hypothèse modifiable · noir = formule · fond jaune = hypothèse clé. "
        "Montants en € HT. Aucune donnée ne provient d'un client réel : tout est hypothèse à valider sur le terrain.").font = NOTE
row[0] = 4
r = hyp("scen", "Scénario actif (1 = pessimiste · 2 = central · 3 = optimiste)", SCENARIO, '0', key=True)

section(hy, "1. OFFRES ET PRIX (abonnement mensuel, engagement 12 mois)")
hyp("p1", "Offre « Caveau » — prix mensuel", 59, note="WinePass : QR dégustation, fiche vins, contacts, relances FR")
hyp("p2", "Offre « Export » — prix mensuel", 129, note="+ réachat international, prix rendu partenaires, 4 langues, suivi")
hyp("p3", "Offre « Signature » — prix mensuel", 249, note="+ multi-sites, intégrations e-shop/caisse, coaching trimestriel")
hyp("m1", "Mix des nouveaux clients — Caveau", 0.40, PCT)
hyp("m2", "Mix des nouveaux clients — Export", 0.45, PCT)
hyp("m3", "Mix des nouveaux clients — Signature", 0.15, PCT)
hyp("pavg", "Prix mensuel moyen pondéré (ARPA abonnement)",
    f"={H['p1']}*{H['m1']}+{H['p2']}*{H['m2']}+{H['p3']}*{H['m3']}", '#,##0.00', key=True)
hyp("setup", "Frais de mise en service (par nouveau client)", 390, note="Paramétrage, import catalogue, formation équipe 1 h, kit QR")
hyp("pilots", "Nombre de domaines pilotes (mise en service offerte + 3 mois gratuits)", 10, NB)
hyp("freem", "Mois gratuits accordés aux pilotes", 3, NB)

section(hy, "2. REVENUS COMPLÉMENTAIRES")
hyp("com", "Commission sur la valeur vin des réachats passés par la plateforme", 0.02, PCT)
hyp("basket", "Panier vin moyen d'un réachat international (€ TTC encaissés par le domaine)", 250)
hyp("elig", "Part des clients ayant le réachat international (Export + Signature)", f"={H['m2']}+{H['m3']}", PCT)
hyp_years("orders", "Réachats par client éligible et par an (montée en charge du fichier)", [8, 14, 18], NB,
          "Hypothèse prudente, à mesurer chez les pilotes")
hyp("attach", "Part des nouveaux clients achetant un pack DomaineDigital", 0.30, PCT)
hyp("ticket", "Prix moyen d'un pack DomaineDigital", 750)
hyp("svccost", "Coût des freelances pour les packs (en % du prix)", 0.60, PCT)

section(hy, "3. ACQUISITION ET RÉTENTION (pilotées par le scénario)", 6)
r = row[0]
for i, t in enumerate(["", "", "Pessimiste", "Central", "Optimiste", "→ retenu"]):
    hy.cell(row=r, column=1 + i, value=t).font = BOLD
row[0] = r + 1
SC = {}
def scen_row(name, label, vals, fmt):
    r = row[0]
    hy.cell(row=r, column=1, value=label).font = BLACK
    for i, v in enumerate(vals):
        c = hy.cell(row=r, column=3 + i, value=v)
        c.font = BLUE
        c.number_format = fmt
    c = hy.cell(row=r, column=6, value=f"=CHOOSE({H['scen']},C{r},D{r},E{r})")
    c.font = BOLD
    c.number_format = fmt
    c.fill = KEY_FILL
    SC[name] = f"Hypothèses!$F${r}"
    row[0] = r + 1
scen_row("new27", "Nouveaux domaines signés — 2027", [25, 45, 65], NB)
scen_row("new28", "Nouveaux domaines signés — 2028", [80, 125, 180], NB)
scen_row("new29", "Nouveaux domaines signés — 2029", [130, 200, 290], NB)
scen_row("churn", "Churn mensuel (résiliations / clients actifs)", [0.025, 0.015, 0.010], '0.00%')

section(hy, "4. SAISONNALITÉ DES SIGNATURES (part annuelle signée chaque mois)", 14)
r = row[0]
months = ["janv", "févr", "mars", "avr", "mai", "juin", "juil", "août", "sept", "oct", "nov", "déc"]
for i, m in enumerate(months):
    hy.cell(row=r, column=3 + i, value=m).font = BOLD
row[0] = r + 1
r27 = row[0]
hy.cell(row=r27, column=1, value="2027 (MVP disponible en mars : pas de signature avant)").font = BLACK
for i, v in enumerate([0, 0, .07, .11, .11, .07, .04, .02, .09, .13, .18, .18]):
    c = hy.cell(row=r27, column=3 + i, value=v); c.font = BLUE; c.number_format = '0%'
hy.cell(row=r27, column=15, value=f"=SUM(C{r27}:N{r27})").number_format = '0%'
r2829 = r27 + 1
hy.cell(row=r2829, column=1, value="2028-2029 (vente en basse saison caveau : nov. → avril)").font = BLACK
for i, v in enumerate([.10, .12, .14, .13, .08, .05, .03, .02, .05, .08, .10, .10]):
    c = hy.cell(row=r2829, column=3 + i, value=v); c.font = BLUE; c.number_format = '0%'
hy.cell(row=r2829, column=15, value=f"=SUM(C{r2829}:N{r2829})").number_format = '0%'
row[0] = r2829 + 1

section(hy, "5. ÉQUIPE (coût employeur mensuel chargé, € — 0 = absent)", 7)
r = row[0]
for i, t in enumerate(["Rôle", "Entrée (mois n°)", "2027", "2028", "2029"]):
    hy.cell(row=r, column=1 + i if i == 0 else 1 + i, value=t).font = BOLD
row[0] = r + 1
TEAM = [
    ("CEO — fondatrice (vente, partenariats, finance)", 3, [2300, 3600, 4600]),
    ("CTO — cofondateur associé (produit, tech)", 3, [2300, 3600, 4600]),
    ("Business developer terrain n°1", 11, [3900, 4000, 4100]),
    ("Alternant(e) marketing digital", 11, [950, 950, 950]),
    ("Développeur(se) full-stack", 15, [0, 4800, 4950]),
    ("Business developer terrain n°2", 23, [0, 4000, 4100]),
    ("Customer success / onboarding", 27, [0, 0, 3500]),
]
TEAMROWS = []
for label, start, costs in TEAM:
    r = row[0]
    hy.cell(row=r, column=1, value=label).font = BLACK
    c = hy.cell(row=r, column=2, value=start); c.font = BLUE
    for i, v in enumerate(costs):
        c = hy.cell(row=r, column=3 + i, value=v); c.font = BLUE; c.number_format = EUR
    TEAMROWS.append(r)
    H[f"start{len(TEAMROWS)-1}"] = f"Hypothèses!$B${r}"
    row[0] = r + 1
hy.cell(row=row[0], column=1, value="Mois n° : 1 = nov. 2026, 3 = janv. 2027, 11 = sept. 2027, 15 = janv. 2028, 23 = sept. 2028, 27 = janv. 2029. "
        "Fondateurs non rémunérés pendant l'ante-création. Variable commercial inclus dans le coût BizDev.").font = NOTE
row[0] += 1

section(hy, "6. CHARGES D'EXPLOITATION", 7)
r = row[0]
for i, t in enumerate(["", "", "2027", "2028", "2029"]):
    hy.cell(row=r, column=1 + i, value=t).font = BOLD
row[0] = r + 1
hyp_years("infra_fix", "Hébergement UE, monitoring, licences techniques — fixe mensuel", [250, 400, 600])
hyp_years("infra_var", "Coût variable par client actif et par mois (hébergement, e-mails/SMS, traduction API)", [5, 5, 5])
hyp_years("mkt_dig", "Marketing digital (LinkedIn, contenus, SEO) — mensuel", [300, 700, 1200])
hyp_years("travel", "Déplacements terrain — mensuel", [700, 1500, 2500])
hyp_years("salon_feb", "Salons — février (Wine Paris / Vinexpo Paris)", [3000, 6000, 10000])
hyp_years("salon_jun", "Salons — juin (salons régionaux, rencontres œnotourisme)", [0, 3000, 5000])
hyp_years("salon_nov", "Salons — novembre (Sitevi / salons des vignerons)", [2000, 6000, 10000])
hyp_years("overhead", "Frais généraux mensuels (comptable, assurances RC/cyber, banque, outils internes)", [550, 800, 1100])
hyp_years("office", "Coworking — mensuel", [200, 500, 900])
hyp_years("legal", "Juridique annuel (CGU/CGV, RGPD, contrats partenaires) — janvier", [0, 2000, 2500])
hyp("referral", "Prime d'apport versée aux prescripteurs, par client apporté", 100)
hyp("refshare", "Part des nouveaux clients apportés par des prescripteurs", 0.30, PCT)
hyp("free_mvp", "Développement MVP sous-traité (freelance) — par mois, déc. 2026 → avril 2027", 5000)
hyp("free_mnt", "Maintenance freelance — par mois, mai → déc. 2027", 800)
hyp("ante_legal", "Ante-création : statuts, CGU/CGV, RGPD, marque INPI, création (nov. 2026)", 5500)
hyp("ante_site", "Ante-création : site vitrine, identité visuelle, maquettes (déc. 2026)", 2500)
hyp("capex0", "Investissement initial : 2 ordinateurs + 2 tablettes de démonstration (déc. 2026)", 3800)
hyp("capexhire", "Équipement par recrutement (ordinateur)", 1400)
hyp("amort", "Durée d'amortissement du matériel (mois)", 36, NB)

section(hy, "7. FINANCEMENT", 7)
hyp("apport", "Apport en capital des fondateurs (nov. 2026)", 10000)
hyp("honneur", "Prêts d'honneur des fondateurs apportés en compte courant (déc. 2026)", 40000,
    note="2 × 20 k€, prêts personnels à 0 %, remboursés par les fondateurs")
hyp("bft1", "Subvention Bourse French Tech — 1re tranche (mars 2027)", 21000, note="À confirmer auprès de Bpifrance")
hyp("bft2", "Subvention Bourse French Tech — solde (mars 2028)", 9000)
hyp("bank", "Prêt bancaire (janv. 2027)", 60000)
hyp("rate", "Taux annuel du prêt bancaire", 0.045, PCT)
hyp("dur", "Durée du prêt (mois), remboursement dès février 2027", 60, NB)
hyp("ba", "Levée en fonds propres — business angels (mars 2027) : MONTANT SOLLICITÉ", 120000, key=True)
hyp("seed", "Levée seed éventuelle (non retenue dans le central)", 0)
hyp("is_low", "IS — taux réduit PME (jusqu'à 42 500 € de bénéfice)", 0.15, PCT)
hyp("is_high", "IS — taux normal", 0.25, PCT)

# =========================================================== MENSUEL
ms = wb.create_sheet("Mensuel")
ms.column_dimensions["A"].width = 46
ms.column_dimensions["B"].width = 9
N = 38                     # nov. 2026 → déc. 2029
C0 = 3                     # colonne C = mois 1
def col(m):                # m = 1..38
    return L(C0 + m - 1)
FIRST, LAST = col(1), col(N)
ms.freeze_panes = "C5"
title(ms, 1, "MOTEUR MENSUEL — nov. 2026 → déc. 2029 (mois 1 à 38)")
labels_month = []
for m in range(1, N + 1):
    y = 2026 + (m + 9) // 12
    mo = (m + 9) % 12 + 1
    labels_month.append((y, mo))
    ms.cell(row=3, column=C0 + m - 1, value=f"{months[mo-1]} {str(y)[2:]}").font = BOLD
    ms.cell(row=4, column=C0 + m - 1, value=y).font = NOTE
    ms.column_dimensions[col(m)].width = 10
ms.cell(row=3, column=1, value="Mois").font = BOLD
ms.cell(row=4, column=1, value="Exercice").font = NOTE

R = {}
cur = [6]
def line(name, label, fn, fmt=EUR, bold=False, fill=None):
    r = cur[0]
    R[name] = r
    ms.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for m in range(1, N + 1):
        y, mo = labels_month[m - 1]
        v = fn(m, y, mo, r)
        c = ms.cell(row=r, column=C0 + m - 1, value=v)
        c.number_format = fmt
        c.font = BOLD if bold else BLACK
        if fill:
            c.fill = fill
    if fill:
        ms.cell(row=r, column=1).fill = fill
    R[name] = r
    cur[0] = r + 1
def blank(label=None):
    r = cur[0]
    if label:
        for cc in range(1, C0 + N):
            ms.cell(row=r, column=cc).fill = SEC_FILL
        ms.cell(row=r, column=1, value=label).font = SEC
    cur[0] = r + 1
def ref(name, m):
    return f"{col(m)}{R[name]}"
def prev(name, m):
    return f"{col(m-1)}{R[name]}" if m > 1 else "0"
def ycol(y):               # colonne C/D/E de l'onglet Hypothèses
    return {2027: "C", 2028: "D", 2029: "E"}.get(y, "C")

blank("CLIENTS")
def newc(m, y, mo, r):
    if y == 2026:
        return 0
    srow = r27 if y == 2027 else r2829
    total = SC["new27"] if y == 2027 else SC["new28"] if y == 2028 else SC["new29"]
    return f"={total}*Hypothèses!{L(2 + mo)}${srow}"
line("new", "Nouveaux domaines signés", newc, NB1)
line("cumnew", "Cumul des signatures", lambda m, y, mo, r: f"={prev('cumnew', m)}+{ref('new', m)}", NB1)
line("active", "Domaines actifs (fin de mois)", lambda m, y, mo, r: f"={prev('active', m)}*(1-{SC['churn']})+{ref('new', m)}", NB1, bold=True)
line("churned", "Résiliations", lambda m, y, mo, r: f"={prev('active', m)}*{SC['churn']}", NB1)
line("pilotnew", "dont pilotes signés ce mois", lambda m, y, mo, r: f"=MIN({ref('new', m)},MAX(0,{H['pilots']}-{prev('cumnew', m)}))", NB1)
line("setupn", "Mises en service facturées", lambda m, y, mo, r: f"={ref('new', m)}-{ref('pilotnew', m)}", NB1)

blank("REVENUS (€ HT)")
line("mrr", "MRR — revenu mensuel récurrent (abonnements)", lambda m, y, mo, r: f"={ref('active', m)}*{H['pavg']}", EUR, bold=True)
def disc(m, y, mo, r):
    terms = [ref('pilotnew', k) for k in range(max(1, m - 2), m + 1)]
    return f"=-{H['pavg']}*({'+'.join(terms)})"
line("pilotdisc", "Mois offerts aux pilotes", disc)
line("rsub", "CA abonnements facturé", lambda m, y, mo, r: f"={ref('mrr', m)}+{ref('pilotdisc', m)}")
line("rsetup", "CA mises en service", lambda m, y, mo, r: f"={ref('setupn', m)}*{H['setup']}")
line("rcom", "CA commissions réachats", lambda m, y, mo, r:
     "0" if y == 2026 else f"={ref('active', m)}*{H['elig']}*Hypothèses!${ycol(y)}${H['orders']}/12*{H['basket']}*{H['com']}")
line("rsvc", "CA packs DomaineDigital", lambda m, y, mo, r: f"={ref('new', m)}*{H['attach']}*{H['ticket']}")
line("ca", "CHIFFRE D'AFFAIRES TOTAL", lambda m, y, mo, r: f"=SUM({col(m)}{R['rsub']}:{col(m)}{R['rsvc']})", EUR, bold=True, fill=TOT_FILL)

blank("COÛTS DIRECTS")
line("cvar", "Coûts variables d'infrastructure", lambda m, y, mo, r:
     "0" if y == 2026 else f"={ref('active', m)}*Hypothèses!${ycol(y)}${H['infra_var']}")
line("csvc", "Freelances packs DomaineDigital", lambda m, y, mo, r: f"={ref('rsvc', m)}*{H['svccost']}")
line("mb", "MARGE BRUTE", lambda m, y, mo, r: f"={ref('ca', m)}-{ref('cvar', m)}-{ref('csvc', m)}", EUR, bold=True, fill=TOT_FILL)

blank("CHARGES D'EXPLOITATION")
for i, tr in enumerate(TEAMROWS):
    def sal(m, y, mo, r, tr=tr):
        if y == 2026:
            return "0"
        return f"=IF({m}>=Hypothèses!$B${tr},Hypothèses!${ycol(y)}${tr},0)"
    line(f"sal{i}", f"Salaire chargé — {TEAM[i][0]}", sal)
line("salaries", "Total salaires et charges sociales", lambda m, y, mo, r:
     f"=SUM({col(m)}{R['sal0']}:{col(m)}{R['sal%d' % (len(TEAM)-1)]})", EUR, bold=True)
def freel(m, y, mo, r):
    if (y == 2026 and mo == 12) or (y == 2027 and mo <= 4):
        return f"={H['free_mvp']}"
    if y == 2027:
        return f"={H['free_mnt']}"
    return "0"
line("freelance", "Développement sous-traité (MVP puis maintenance)", freel)
line("infra", "Hébergement et licences techniques (fixe)", lambda m, y, mo, r:
     "0" if y == 2026 else f"=Hypothèses!${ycol(y)}${H['infra_fix']}")
def mkt(m, y, mo, r):
    if y == 2026:
        return "0"
    yc = ycol(y)
    s = f"=Hypothèses!${yc}${H['mkt_dig']}+Hypothèses!${yc}${H['travel']}"
    if mo == 2:
        s += f"+Hypothèses!${yc}${H['salon_feb']}"
    if mo == 6:
        s += f"+Hypothèses!${yc}${H['salon_jun']}"
    if mo == 11:
        s += f"+Hypothèses!${yc}${H['salon_nov']}"
    s += f"+{ref('new', m)}*{H['refshare']}*{H['referral']}"
    return s
line("mkt", "Marketing, salons, déplacements, primes d'apport", mkt)
def ovh(m, y, mo, r):
    if y == 2026:
        return f"={H['ante_legal']}" if mo == 11 else f"={H['ante_site']}"
    yc = ycol(y)
    s = f"=Hypothèses!${yc}${H['overhead']}+Hypothèses!${yc}${H['office']}"
    if mo == 1:
        s += f"+Hypothèses!${yc}${H['legal']}"
    return s
line("ovh", "Frais généraux (dont ante-création)", ovh)
line("opex", "TOTAL CHARGES D'EXPLOITATION", lambda m, y, mo, r:
     f"={ref('salaries', m)}+{ref('freelance', m)}+{ref('infra', m)}+{ref('mkt', m)}+{ref('ovh', m)}", EUR, bold=True)
line("ebitda", "EBE (EBITDA)", lambda m, y, mo, r: f"={ref('mb', m)}-{ref('opex', m)}", EUR, bold=True, fill=TOT_FILL)
line("ebitda3", "EBE glissant 3 mois (repère du point mort)", lambda m, y, mo, r:
     f"=SUM({col(max(1, m-2))}{R['ebitda']}:{col(m)}{R['ebitda']})")
line("pmflag", "Point mort atteint (EBE glissant 3 mois > 0 et le reste)", lambda m, y, mo, r:
     f"=IF(COUNTIF({col(m)}{R['ebitda3']}:${LAST}${R['ebitda3']},\"<=0\")=0,1,0)", '0')

blank("INVESTISSEMENTS ET AMORTISSEMENTS")
def capex(m, y, mo, r):
    parts = []
    if m == 2:
        parts.append(H['capex0'])
    for tr in TEAMROWS[2:]:
        parts.append(f"IF({m}=Hypothèses!$B${tr},{H['capexhire']},0)")
    return "=" + "+".join(parts)
line("capex", "Investissements matériels", capex)
line("dep", "Dotations aux amortissements", lambda m, y, mo, r:
     f"=SUM({col(max(1, m-35))}{R['capex']}:{col(m)}{R['capex']})/{H['amort']}")

blank("FINANCEMENT")
line("pmt", "Échéance du prêt bancaire", lambda m, y, mo, r:
     f"=IF(AND({m}>=4,{m}<4+{H['dur']}),-PMT({H['rate']}/12,{H['dur']},{H['bank']}),0)")
line("debt", "Capital restant dû (fin de mois)", lambda m, y, mo, r:
     (f"=IF({m}=3,{H['bank']},0)" if m <= 3 else
      f"={prev('debt', m)}-({ref('pmt', m)}-{prev('debt', m)}*{H['rate']}/12)"))
line("interest", "Intérêts d'emprunt", lambda m, y, mo, r:
     "0" if m <= 3 else f"={prev('debt', m)}*{H['rate']}/12")
line("principal", "Remboursement du capital", lambda m, y, mo, r: f"={ref('pmt', m)}-{ref('interest', m)}")
def inflow(m, y, mo, r):
    parts = []
    if m == 1: parts.append(H['apport'])
    if m == 2: parts.append(H["honneur"])
    if m == 3: parts.append(H["bank"])
    if m == 5: parts += [H['bft1'], H['ba']]
    if m == 17: parts.append(H['bft2'])
    return "=" + "+".join(parts) if parts else "0"
line("inflow", "Apports, prêts, subventions, levée", inflow)
line("grant", "dont subvention (produit exceptionnel en CR)", lambda m, y, mo, r:
     f"={H['bft1']}" if m == 5 else f"={H['bft2']}" if m == 17 else "0")

blank("RÉSULTAT ET TRÉSORERIE")
line("ebit", "Résultat d'exploitation", lambda m, y, mo, r: f"={ref('ebitda', m)}-{ref('dep', m)}", EUR, bold=True)
line("rcai", "Résultat avant impôt (y c. subvention)", lambda m, y, mo, r:
     f"={ref('ebit', m)}-{ref('interest', m)}+{ref('grant', m)}")
line("tax", "Impôt sur les sociétés (calculé en décembre, onglet Synthèse)", lambda m, y, mo, r:
     (f"=Synthèse!{ {2027:'C',2028:'D',2029:'E'}[y] }$IS_ROW" if mo == 12 and y >= 2027 else "0"))
line("net", "Résultat net", lambda m, y, mo, r: f"={ref('rcai', m)}-{ref('tax', m)}", EUR, bold=True)
line("cashflow", "Flux de trésorerie du mois", lambda m, y, mo, r:
     f"={ref('ebitda', m)}-{ref('capex', m)}-{ref('pmt', m)}-{ref('tax', m)}+{ref('inflow', m)}")
line("cash", "TRÉSORERIE FIN DE MOIS", lambda m, y, mo, r: f"={prev('cash', m)}+{ref('cashflow', m)}", EUR, bold=True, fill=TOT_FILL)

# =========================================================== SYNTHÈSE
sy = wb.create_sheet("Synthèse")
sy.column_dimensions["A"].width = 52
for c in "BCDEF":
    sy.column_dimensions[c].width = 15
title(sy, 1, "SYNTHÈSE — compte de résultat, plan de financement, KPI, seuil de rentabilité")
YRS = [(2026, "B", "Ante-création\n(nov.-déc. 26)"), (2027, "C", "2027"), (2028, "D", "2028"), (2029, "E", "2029")]
S = {}
srow = [3]
def shead(text):
    r = srow[0] + 1
    for cc in range(1, 6):
        sy.cell(row=r, column=cc).fill = SEC_FILL
    sy.cell(row=r, column=1, value=text).font = SEC
    for y, c, lab in YRS:
        x = sy[f"{c}{r}"]; x.value = lab; x.font = SEC
        x.alignment = Alignment(wrap_text=True, horizontal="center")
    srow[0] = r + 1
def sline(name, label, fn, fmt=EUR, bold=False, fill=None, years=YRS):
    r = srow[0]
    sy.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for y, c, _ in years:
        v = fn(y, c)
        x = sy[f"{c}{r}"]; x.value = v; x.number_format = fmt
        x.font = BOLD if bold else BLACK
        if fill: x.fill = fill
    if fill: sy.cell(row=r, column=1).fill = fill
    S[name] = r
    srow[0] = r + 1
def ysum(name):
    return lambda y, c: f"=SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R[name]}:${LAST}${R[name]})"
def dec(name):
    # valeur de décembre de l'exercice
    return lambda y, c: f"=Mensuel!{col({2026: 2, 2027: 14, 2028: 26, 2029: 38}[y])}{R[name]}"

shead("COMPTE DE RÉSULTAT (€ HT)")
sline("ca_sub", "CA abonnements", ysum("rsub"))
sline("ca_set", "CA mises en service", ysum("rsetup"))
sline("ca_com", "CA commissions réachats", ysum("rcom"))
sline("ca_svc", "CA packs DomaineDigital", ysum("rsvc"))
sline("ca", "Chiffre d'affaires total", ysum("ca"), bold=True, fill=TOT_FILL)
sline("cd", "Coûts directs (infrastructure variable + freelances)", lambda y, c: f"=SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['cvar']}:${LAST}${R['cvar']})+SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['csvc']}:${LAST}${R['csvc']})")
sline("mb", "Marge brute", ysum("mb"), bold=True)
sline("mbpct", "Taux de marge brute", lambda y, c: f"=IF({c}{S['ca']}=0,0,{c}{S['mb']}/{c}{S['ca']})", PCT)
sline("sal", "Salaires et charges sociales", ysum("salaries"))
sline("tech", "Développement sous-traité + hébergement", lambda y, c: f"=SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['freelance']}:${LAST}${R['freelance']})+SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['infra']}:${LAST}${R['infra']})")
sline("mkt", "Marketing, salons, déplacements", ysum("mkt"))
sline("ovh", "Frais généraux", ysum("ovh"))
sline("ebitda", "EBE (EBITDA)", ysum("ebitda"), bold=True, fill=TOT_FILL)
sline("dep", "Amortissements", ysum("dep"))
sline("int", "Charges financières", ysum("interest"))
sline("grant", "Subvention d'exploitation (Bourse French Tech)", ysum("grant"))
sline("rcai", "Résultat avant impôt", ysum("rcai"), bold=True)
# IS avec report déficitaire (la colonne B = ante-création appartient à l'exercice 2027 : on l'agrège à 2027)
sline("base", "Base imposable cumulée (report des déficits)", lambda y, c:
      "0" if y == 2026 else (f"=B{srow[0]-1}+C{srow[0]-1}" if y == 2027 else f"=MIN(0,{chr(ord(c)-1)}{srow[0]})+{c}{srow[0]-1}"))
sline("is", "Impôt sur les sociétés", lambda y, c:
      "0" if y == 2026 else f"=IF({c}{S['base']}<=0,0,MIN({c}{S['base']},42500)*{H['is_low']}+MAX(0,{c}{S['base']}-42500)*{H['is_high']})")
sline("net", "Résultat net", lambda y, c: f"={c}{S['rcai']}-{c}{S['is']}", bold=True, fill=TOT_FILL)

shead("PLAN DE FINANCEMENT (€)")
sline("e_ebitda", "Capacité d'autofinancement (EBE − IS − intérêts)", lambda y, c: f"={c}{S['ebitda']}-{c}{S['is']}-{c}{S['int']}+{c}{S['grant']}")
sline("e_apport", "Apport en capital des fondateurs", lambda y, c: f"={H['apport']}" if y == 2026 else "0")
sline("e_cc", "Comptes courants (prêts d'honneur des fondateurs)", lambda y, c: f"={H['honneur']}" if y == 2026 else "0")
sline("e_ba", "Levée de fonds propres (business angels)", lambda y, c: f"={H['ba']}" if y == 2027 else "0")
sline("e_bank", "Emprunt bancaire", lambda y, c: f"={H['bank']}" if y == 2027 else "0")
sline("e_tot", "Total ressources (hors subvention, incluse dans la CAF)", lambda y, c: f"=SUM({c}{S['e_apport']}:{c}{S['e_bank']})", bold=True)
sline("u_capex", "Investissements matériels", ysum("capex"))
sline("u_princ", "Remboursement du capital emprunté", ysum("principal"))
sline("u_tot", "Total emplois", lambda y, c: f"={c}{S['u_capex']}+{c}{S['u_princ']}", bold=True)
sline("var", "Variation de trésorerie", lambda y, c: f"={c}{S['e_ebitda']}+{c}{S['e_tot']}-{c}{S['u_tot']}", bold=True)
sline("cash_end", "Trésorerie de fin de période", dec("cash"), bold=True, fill=TOT_FILL)
sline("cash_min", "Trésorerie mensuelle minimale de la période", lambda y, c:
      f"=MIN(Mensuel!{col({2026:1,2027:3,2028:15,2029:27}[y])}{R['cash']}:Mensuel!{col({2026:2,2027:14,2028:26,2029:38}[y])}{R['cash']})")

shead("KPI SaaS")
K = YRS[1:]
sline("k_new", "Nouveaux domaines signés", ysum("new"), NB, years=K)
sline("k_act", "Domaines actifs au 31/12", dec("active"), NB, years=K)
sline("k_mrr", "MRR de décembre", dec("mrr"), EUR, years=K)
sline("k_arr", "ARR (MRR de décembre × 12)", lambda y, c: f"={c}{S['k_mrr']}*12", EUR, bold=True, years=K)
sline("k_churn", "Churn annuel équivalent", lambda y, c: f"=1-(1-{SC['churn']})^12", PCT, years=K)
sline("k_arpa", "Revenu récurrent annuel moyen par domaine (abonnement + commission)", lambda y, c:
      f"=({c}{S['ca_sub']}+{c}{S['ca_com']})/AVERAGE(Mensuel!{col({2027:3,2028:15,2029:27}[y])}{R['active']}:Mensuel!{col({2027:14,2028:26,2029:38}[y])}{R['active']})", EUR, years=K)
sline("k_sm", "Dépenses commerciales et marketing (BizDev + 50 % CEO + marketing + alternant)", lambda y, c:
      f"=SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['mkt']}:${LAST}${R['mkt']})"
      f"+SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['sal2']}:${LAST}${R['sal2']})"
      f"+SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['sal3']}:${LAST}${R['sal3']})"
      f"+SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['sal5']}:${LAST}${R['sal5']})"
      f"+0.5*SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['sal0']}:${LAST}${R['sal0']})", EUR, years=K)
sline("k_cac", "CAC — coût d'acquisition par domaine", lambda y, c: f"={c}{S['k_sm']}/{c}{S['k_new']}", EUR, bold=True, years=K)
sline("k_gm", "Marge brute sur revenus récurrents (hors packs)", lambda y, c:
      f"=1-SUMIF(Mensuel!${FIRST}$4:${LAST}$4,{y},Mensuel!${FIRST}${R['cvar']}:${LAST}${R['cvar']})/({c}{S['ca_sub']}+{c}{S['ca_com']})", PCT, years=K)
sline("k_ltv", "LTV = ARPA × marge brute / churn annuel", lambda y, c: f"={c}{S['k_arpa']}*{c}{S['k_gm']}/{c}{S['k_churn']}", EUR, bold=True, years=K)
sline("k_ratio", "Ratio LTV / CAC", lambda y, c: f"={c}{S['k_ltv']}/{c}{S['k_cac']}", '0.0', years=K)
sline("k_payback", "Délai de récupération du CAC (mois)", lambda y, c: f"={c}{S['k_cac']}/({c}{S['k_arpa']}*{c}{S['k_gm']}/12)", '0.0', years=K)

shead("SEUIL DE RENTABILITÉ")
sline("b_fixed", "Charges fixes annuelles (opex + amortissements + intérêts − subvention)", lambda y, c:
      f"={c}{S['sal']}+{c}{S['tech']}+{c}{S['mkt']}+{c}{S['ovh']}+{c}{S['dep']}+{c}{S['int']}-{c}{S['grant']}", EUR, years=K)
sline("b_contrib", "Contribution annuelle d'un domaine actif (abonnement + commission − coût variable)", lambda y, c:
      f"=12*{H['pavg']}+{H['elig']}*Hypothèses!{ {2027:'C',2028:'D',2029:'E'}[y] }${H['orders']}*{H['basket']}*{H['com']}-12*Hypothèses!{ {2027:'C',2028:'D',2029:'E'}[y] }${H['infra_var']}", EUR, years=K)
sline("b_units", "Seuil : nombre de domaines actifs nécessaires (moyenne annuelle)", lambda y, c: f"={c}{S['b_fixed']}/{c}{S['b_contrib']}", NB, bold=True, years=K)
sline("b_mrr", "Seuil exprimé en MRR", lambda y, c: f"={c}{S['b_units']}*{H['pavg']}", EUR, years=K)
r_pm = srow[0] + 1
sy.cell(row=r_pm, column=1, value="Point mort : mois à partir duquel l'EBE glissant sur 3 mois reste positif (n° de mois, 1 = nov. 2026)").font = BOLD
sy.cell(row=r_pm, column=2, value=f"=IFERROR(MATCH(1,Mensuel!{FIRST}{R['pmflag']}:{LAST}{R['pmflag']},0),\"au-delà de 2029\")").font = BOLD
S["pointmort"] = r_pm
sy.cell(row=r_pm + 1, column=1, value="Trésorerie minimale sur toute la période (€)").font = BOLD
sy.cell(row=r_pm + 1, column=2, value=f"=MIN(Mensuel!{FIRST}{R['cash']}:{LAST}{R['cash']})").number_format = EUR
S["cashlow"] = r_pm + 1

# remplace le marqueur IS_ROW dans l'onglet Mensuel
for m in range(1, N + 1):
    c = ms.cell(row=R["tax"], column=C0 + m - 1)
    if isinstance(c.value, str) and "IS_ROW" in c.value:
        c.value = c.value.replace("IS_ROW", str(S["is"]))

wb.save(OUT)
import json
json.dump({"H": H, "SC": SC, "R": R, "S": S, "N": N}, open(OUT.replace(".xlsx", "_map.json"), "w"), ensure_ascii=False)
print("ok", OUT)

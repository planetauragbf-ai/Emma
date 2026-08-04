# -*- coding: utf-8 -*-
"""Génère annexes/Tableaux_De_Bord.xlsx — tableaux de bord de contrôle de l'activité.

Exigence CDC : « les tableaux de bord correspondent aux objectifs fixés ».
Les objectifs sont donc FIGÉS depuis annexes/Modele_Financier_BP.xlsx (valeurs
recalculées) — pas de lien inter-fichiers (fragile), mais une note de version.
Colonnes « Réalisé » à remplir en exploitation ; « Écart » en formule.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "/home/user/Emma/annexes/Modele_Financier_BP.xlsx"
OUT = "/home/user/Emma/annexes/Tableaux_De_Bord.xlsx"

ARIAL = "Arial"
BLACK = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=13, bold=True)
NOTE = Font(name=ARIAL, size=9, italic=True, color="666666")
EX = Font(name=ARIAL, size=10, italic=True, color="808080")
SECTION = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="2F5B69")
HDR_FILL = PatternFill("solid", fgColor="DCE6EA")
IN_FILL = PatternFill("solid", fgColor="FFF2CC")   # cellules à remplir
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
EUR = '#,##0;(#,##0);"-"'
PCT = '0.0%'
NB = '#,##0'

src = load_workbook(SRC, data_only=True)
tq = src["Trimestriel"]
QC = [get_column_letter(2 + i) for i in range(12)]

def row_vals(r):
    return [tq[f"{c}{r}"].value or 0 for c in QC]

adds = row_vals(4)
clients_fin = row_vals(6)
sal_fin = row_vals(8)
sal_moy = row_vals(9)
pilotes = row_vals(14)
prix_moy = row_vals(12)
mrr_fin = [s * p for s, p in zip(sal_fin, prix_moy)]
arr_fin = [m * 12 for m in mrr_fin]
mb_pct = row_vals(26)
ebitda = row_vals(36)
treso = row_vals(43)
packs_cible = [tq[f"{c}16"].value / 110 if tq[f"{c}16"].value else 0 for c in QC]
demos = [p / 0.4 for p in pilotes]
rdv = [d / 0.5 for d in demos]
comptes = [r / 0.25 for r in rdv]

wb = Workbook()

# ---------------------------------------------------------------- Mode d'emploi
me = wb.active
me.title = "Mode d'emploi"
me.column_dimensions["A"].width = 112
me["A1"] = "TABLEAUX DE BORD DE CONTRÔLE DE L'ACTIVITÉ"
me["A1"].font = TITLE
texte = [
    "",
    "Trois tableaux : Commercial (l'entonnoir et le portefeuille) · Adoption (la valeur d'usage — notre hypothèse",
    "la plus risquée) · Finance (l'économie du projet). Chaque indicateur a trois lignes :",
    "   OBJECTIF — figé depuis le modèle financier (Modele_Financier_BP.xlsx, version du 04/08/2026).",
    "   RÉALISÉ — cellules jaunes, à saisir chaque trimestre (la colonne « Exemple », grisée, montre le format attendu).",
    "   ÉCART — calculé automatiquement (réalisé / objectif − 1) ; s'affiche dès qu'un réalisé est saisi.",
    "",
    "Règles de pilotage (proposées) :",
    "   · Un écart < −20 % sur deux trimestres consécutifs déclenche un plan d'action documenté.",
    "   · L'activation à 3 mois sous 12 % (seuil d'alerte) déclenche le plan de relance adoption",
    "     (sponsoring direction, onboarding managers, campagne interne) avant tout nouveau déploiement.",
    "   · La trésorerie sous 6 mois de burn déclenche la préparation d'un financement relais.",
    "",
    "Pourquoi ces indicateurs : ils correspondent un à un aux objectifs du business plan (critère explicite",
    "de la grille de soutenance : « les tableaux de bord correspondent aux objectifs fixés »).",
]
for i, t in enumerate(texte, start=2):
    me[f"A{i}"] = t
    me[f"A{i}"].font = NOTE if t.startswith("   ") else BLACK

# ---------------------------------------------------------------- feuilles KPI
def kpi_sheet(name, blocs):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 11
    for i in range(3, 15):
        ws.column_dimensions[get_column_letter(i)].width = 10.5
    ws["A1"] = f"TABLEAU DE BORD — {name.upper()}"
    ws["A1"].font = TITLE
    ws["A2"] = "Objectifs figés depuis le modèle financier (04/08/2026). Jaune = à saisir. Colonne Exemple = format attendu."
    ws["A2"].font = NOTE
    ws["B3"] = "Exemple"
    ws["B3"].font = EX
    ws["B3"].fill = EX_FILL
    for i, c in enumerate(range(3, 15)):
        cell = ws.cell(row=3, column=c, value=f"T{i+1}")
        cell.font = BOLD
        cell.fill = HDR_FILL
    r = 4
    for bloc_title, indicateurs in blocs:
        ws.cell(row=r, column=1, value=bloc_title).font = SECTION
        for col in range(1, 15):
            ws.cell(row=r, column=col).fill = SEC_FILL
        r += 1
        for label, fmt, objectifs, exemple in indicateurs:
            ws.cell(row=r, column=1, value=f"{label} — objectif").font = BLACK
            for i, v in enumerate(objectifs):
                cc = ws.cell(row=r, column=3 + i, value=v)
                cc.number_format = fmt
            ws.cell(row=r, column=2, value="—").font = EX
            ws.cell(row=r, column=2).fill = EX_FILL
            r += 1
            ws.cell(row=r, column=1, value=f"{label} — réalisé").font = BLACK
            exc = ws.cell(row=r, column=2, value=exemple)
            exc.font = EX
            exc.fill = EX_FILL
            exc.number_format = fmt
            for i in range(12):
                cc = ws.cell(row=r, column=3 + i)
                cc.fill = IN_FILL
                cc.number_format = fmt
            r += 1
            ws.cell(row=r, column=1, value=f"{label} — écart").font = NOTE
            for i in range(12):
                col = get_column_letter(3 + i)
                cc = ws.cell(
                    row=r, column=3 + i,
                    value=f'=IF({col}{r-1}="","",IFERROR({col}{r-1}/{col}{r-2}-1,""))')
                cc.number_format = PCT
                cc.font = NOTE
            r += 1
        r += 1
    return ws

kpi_sheet("Commercial", [
    ("Entonnoir de prospection (cibles issues des taux de passage du BP)", [
        ("Comptes travaillés (cumul trimestre)", NB, [round(x) for x in comptes], 120),
        ("RDV obtenus", NB, [round(x) for x in rdv], 30),
        ("Démos réalisées", NB, [round(x) for x in demos], 15),
        ("Pilotes démarrés", NB, [round(x) for x in pilotes], 6),
        ("Contrats signés (nouveaux clients)", NB, adds, 3),
    ]),
    ("Portefeuille", [
        ("Clients actifs (fin de trimestre)", NB, clients_fin, 4),
        ("Salariés couverts (fin de trimestre)", NB, sal_fin, 1600),
        ("MRR fin de trimestre (€)", EUR, [round(m) for m in mrr_fin], 15000),
    ]),
])

kpi_sheet("Adoption", [
    ("Activation (hypothèse la plus risquée du projet — scénario central 22 %, alerte < 12 %)", [
        ("Taux d'activation à 3 mois (cohortes du trimestre)", PCT, [0.22] * 12, 0.24),
        ("Utilisateurs actifs mensuels (% des comptes activés)", PCT, [0.50] * 12, 0.55),
    ]),
    ("Valeur d'usage", [
        ("NPS salariés (enquête trimestrielle)", NB, [30] * 12, 42),
        ("Actions d'activation réalisées (Cockpit, nb)", NB,
         [round(s * 0.22 * 0.30) for s in sal_moy], 105),
        ("Packs vendus (nb)", NB, [round(p) for p in packs_cible], 4),
        ("Passeports EDUCFI délivrés (nb)", NB,
         [round(s * 0.22 * 0.15) for s in sal_moy], 50),
    ]),
    ("Rétention côté client", [
        ("Clients en pilote convertis (%)", PCT, [0.6] * 12, 0.67),
        ("Renouvellements obtenus / dus (%)", PCT, [1.0] * 12, 1.0),
    ]),
])

kpi_sheet("Finance", [
    ("Revenus", [
        ("ARR fin de trimestre (€)", EUR, [round(a) for a in arr_fin], 180000),
        ("Marge brute (%)", PCT, [round(m, 3) if m else 0 for m in mb_pct], 0.81),
    ]),
    ("Cash", [
        ("EBITDA du trimestre (€)", EUR, [round(e) for e in ebitda], -150000),
        ("Trésorerie fin de trimestre (€)", EUR, [round(t) for t in treso], 1300000),
        ("Runway (mois de charges fixes devant soi)", NB, [""] * 12, 18),
    ]),
    ("Efficacité commerciale (revue semestrielle)", [
        ("CAC (€ / nouveau client)", EUR, [""] * 3 + [8000] + [""] * 3 + [18000] + [""] * 3 + [20600], 7500),
        ("Payback CAC (mois)", NB, [""] * 3 + [3] + [""] * 3 + [5] + [""] * 3 + [4], 3),
    ]),
])

wb.save(OUT)
print("OK —", OUT)

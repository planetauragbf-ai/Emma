"""Génère (V2) livrables/Passavin_Dossier_Investisseur.pptx (30 pages) et livrables/Passavin_Pitch_Soutenance.pptx (20 slides).
Tous les chiffres sont lus dans bp_out.json (scénarios du business plan) : relancer bp_model.py, build_annexes.py puis ce script
après toute modification, puis convertir en PDF (soffice --headless --convert-to pdf).
"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
J = json.load(open(os.path.join(HERE, "bp_out.json")))
S = J["scenarios"]; C = S["central"]; Y = C["years"]; P = C["p"]; U = J["unit"]; MK = J["market"]; LV = J["levers"]
PE, OP, NWL, NSF = S["pessimiste"], S["optimiste"], S["central_sans_marque_blanche"], S["central_sans_frais_logisticien"]
R = C["rows"]
CH = lambda n: os.path.join(HERE, "charts", n)

INK, WINE, TEAL, SAND, GREY, WHITE, GOLD = (RGBColor(0x14, 0x23, 0x3A), RGBColor(0x8C, 0x2F, 0x39), RGBColor(0x2A, 0x7F, 0x87),
                                            RGBColor(0xF4, 0xEF, 0xE6), RGBColor(0x5E, 0x66, 0x73), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xC9, 0xA6, 0x6B))
LIGHT = RGBColor(0xE9, 0xDF, 0xCF)
FONT = "Calibri"


def k(v): return f"{v/1000:,.0f} k€".replace(",", " ").replace("-", "−")
def n(v): return f"{v:,.0f}".replace(",", " ").replace("-", "−")
def src(y, key): return Y[y]["sources"].get(key, 0)
def tot(key): return sum(src(i, key) for i in range(3))


BE = f"M{C['be']}"; BE_OPT = f"M{OP['be']}"
SEED = P["seed_equity"] + P["seed_bpi_loan"]
PRE = P["equity_founder"] + P["honor_loan"] + P["bpi_grant"] + P["bank_loan"]
TOTAL_FUND = PRE + SEED
NEW_TOT = sum(round(y["nouveaux"]) for y in Y)
SAM_TOT = sum(v["sam_3cms"] for v in MK.values()); TAM_TOT = sum(v["tam"] for v in MK.values())
CNT = R[-1]["by_country"]; CMS36 = R[-1]["by_cms"]
H = C["hired"]
AG = tot("agence_silver") + tot("agence_gold")


class Deck:
    def __init__(self, label):
        self.p = Presentation(); self.p.slide_width = Inches(13.333); self.p.slide_height = Inches(7.5)
        self.label = label; self.i = 0

    def _bg(self, s, color):
        f = s.background.fill; f.solid(); f.fore_color.rgb = color

    def text(self, s, x, y, w, h, txt, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.vertical_anchor = anchor; tf.margin_left = tf.margin_right = Inches(0.05)
        lines = txt if isinstance(txt, list) else [txt]
        for j, line in enumerate(lines):
            par = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); par.alignment = align
            r = par.add_run(); r.text = line; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            r.font.italic = italic; r.font.name = FONT
        return tb

    def box(self, s, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid(); b.fill.fore_color.rgb = fill
        if line: b.line.color.rgb = line; b.line.width = Pt(1)
        else: b.line.fill.background()
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = 0.08
        b.shadow.inherit = False
        return b

    def slide(self, title, kicker=None):
        self.i += 1
        s = self.p.slides.add_slide(self.p.slide_layouts[6]); self._bg(s, SAND)
        self.box(s, 0, 0, 0.12, 7.5, WINE, shape=MSO_SHAPE.RECTANGLE)
        if kicker: self.text(s, 0.6, 0.35, 10, 0.35, kicker.upper(), 11, TEAL, True)
        self.text(s, 0.6, 0.65, 12.2, 0.8, title, 26, INK, True)
        self.text(s, 0.6, 7.0, 9, 0.3, f"PASSAVIN · nom de travail provisoire · {self.label} · [Nom Prénom]", 9, GREY)
        self.text(s, 12.0, 7.0, 0.8, 0.3, f"{self.i:02d}", 9, GREY, align=PP_ALIGN.RIGHT)
        return s

    def bullets(self, s, x, y, w, h, items, size=15, color=INK, gap=6):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        for j, it in enumerate(items):
            par = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); par.space_after = Pt(gap)
            bold = it.startswith("**"); it = it.strip("*")
            r = par.add_run(); r.text = ("" if it.startswith(("—", "✔", "✘", "→", "①", "②", "③", "④", "⑤")) else "• ") + it
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT; r.font.bold = bold

    def table(self, s, x, y, w, rows, widths=None, size=11, h_row=0.36, header_fill=INK, bold_first_col=False, h_head=None):
        nr, nc = len(rows), len(rows[0])
        t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h_row * nr)).table
        widths = widths or [w / nc] * nc
        for j, cw in enumerate(widths): t.columns[j].width = Inches(cw)
        for r_, row in enumerate(rows):
            t.rows[r_].height = Inches(h_head if (r_ == 0 and h_head) else h_row)
            for c_, val in enumerate(row):
                cell = t.cell(r_, c_); cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill if r_ == 0 else (WHITE if r_ % 2 else RGBColor(0xFA, 0xF7, 0xF1))
                tf = cell.text_frame; tf.word_wrap = True; par = tf.paragraphs[0]
                txt = str(val); b = txt.startswith("**"); txt = txt.strip("*")
                r = par.add_run(); r.text = txt; r.font.size = Pt(size); r.font.name = FONT
                r.font.color.rgb = WHITE if r_ == 0 else INK; r.font.bold = r_ == 0 or b or (bold_first_col and c_ == 0)
        return t

    def kpis(self, s, x, y, items, w=2.9, h=1.55, gap=0.25, big=30, dark_first=True, lab=12):
        for j, (val, label) in enumerate(items):
            dark = dark_first and j == 0
            self.box(s, x + j * (w + gap), y, w, h, INK if dark else WHITE)
            self.text(s, x + j * (w + gap) + 0.2, y + 0.12, w - 0.4, 0.75, val, big, WHITE if dark else WINE, True)
            self.text(s, x + j * (w + gap) + 0.2, y + 0.85, w - 0.4, h - 0.9, label, lab, LIGHT if dark else GREY)

    def callout(self, s, x, y, w, h, txt, size=15, fill=INK, color=WHITE):
        self.box(s, x, y, w, h, fill)
        self.text(s, x + 0.25, y, w - 0.5, h, txt, size, color, True, anchor=MSO_ANCHOR.MIDDLE)

    def image(self, s, path, x, y, w=None, h=None):
        s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

    def flow(self, s, x, y, steps, w=1.72, h=1.55, gap=0.12, size=11):
        ch = 0.7
        for j, (t1, t2) in enumerate(steps):
            sh = self.box(s, x + j * (w + gap), y, w, ch, INK if j % 2 == 0 else TEAL, shape=MSO_SHAPE.CHEVRON if j else MSO_SHAPE.PENTAGON)
            sh.adjustments[0] = 0.35
            self.text(s, x + j * (w + gap) + (0.3 if j else 0.12), y, w - 0.5, ch, t1, size + 1, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, x + j * (w + gap), y + ch + 0.08, w, h - ch, t2, size, INK)

    def card(self, s, x, y, w, h, title, body, tcolor=TEAL, size=13, tsize=16):
        self.box(s, x, y, w, h, WHITE)
        self.text(s, x + 0.2, y + 0.12, w - 0.4, 0.45, title, tsize, tcolor, True)
        if isinstance(body, list): self.bullets(s, x + 0.2, y + 0.6, w - 0.4, h - 0.7, body, size, gap=3)
        else: self.text(s, x + 0.2, y + 0.6, w - 0.4, h - 0.7, body, size)

    def cover(self, title, subtitle, foot):
        self.i += 1
        s = self.p.slides.add_slide(self.p.slide_layouts[6]); self._bg(s, INK)
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.2), Inches(-1.5), Inches(6), Inches(6)); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x1E, 0x33, 0x52); c.line.fill.background()
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.6), Inches(3.9), Inches(2.6), Inches(2.6)); c.fill.solid(); c.fill.fore_color.rgb = WINE; c.line.fill.background()
        self.text(s, 0.9, 1.9, 9, 1.2, "PASSAVIN", 60, WHITE, True)
        self.box(s, 0.95, 3.15, 1.2, 0.09, WINE, shape=MSO_SHAPE.RECTANGLE)
        self.text(s, 0.9, 3.45, 9.5, 1.0, title, 24, WHITE)
        self.text(s, 0.9, 4.55, 9.5, 0.9, subtitle, 16, LIGHT)
        self.text(s, 0.9, 6.5, 11, 0.5, foot, 11, LIGHT)
        return s

    def save(self, name):
        path = os.path.join(ROOT, "livrables", name); self.p.save(path); return path


# =====================================================================================================
# DOSSIER INVESTISSEUR (30 pages)
# =====================================================================================================
d = Deck("Dossier Bloc 3 · V2")
d.cover("Le connecteur neutre entre la boutique en ligne des domaines européens et les logisticiens vin",
        "Dossier de présentation investisseurs — Bloc 3 Euridis · Mastère Marketing & Digital Business · Version 2",
        "[Nom Prénom] · [Promotion] · Octobre 2026 · Nom de travail provisoire, sous réserve des vérifications INPI/EUIPO")

# 2 — Résumé
s = d.slide("Résumé investisseur", "En une page")
d.kpis(s, 0.6, 1.5, [(n(SAM_TOT), "domaines dans le SAM européen (5 pays, 3 CMS) — ×4 à 7 par rapport à la V1"),
                     (k(Y[2]["arr_fin"]), f"revenu récurrent annualisé fin 2029 (central : {n(Y[2]['clients_fin'])} domaines + {n(Y[2]['wl_domaines'])} en marque blanche)"),
                     (BE, f"point mort mensuel en central (M36 visé, non atteint ; {BE_OPT} en optimiste)"),
                     (k(TOTAL_FUND), f"besoin sur 3 ans : {k(PRE)} de pré-amorçage + seed de {k(SEED)} à M{P['seed_month']}, conditionnée")], w=2.85, lab=11)
d.table(s, 0.6, 3.35, 12.1, [
    ["Problème", "Solution", "Modèle", "Organisation commerciale", "Statut honnête"],
    ["Ressaisie des commandes étrangères, pays non livrables au panier, frais de port faux",
     "Connecteur WooCommerce → Shopify → PrestaShop : éligibilité au checkout, transmission, suivi, multi-logisticiens",
     "Direct : abonnement 49-199 €/mois. Partner : licence de marque blanche 690-2 490 €/mois pour les logisticiens",
     "Dirigeant + 3 sales pays + partner manager + 2 customer success ; 27 agences, 20 apporteurs, 4 logisticiens revendeurs",
     "Aucun client payant. Validation avant le MVP (≈ 3 mois, < 8 k€)"]],
    widths=[2.4, 2.6, 2.5, 2.6, 2.0], size=11, h_row=1.45, h_head=0.4)
d.callout(s, 0.6, 5.85, 12.1, 0.9, f"Nous demandons {k(SEED)} à M{P['seed_month']} ({k(P['seed_equity'])} de fonds propres + {k(P['seed_bpi_loan'])} de prêt d'amorçage BPI), seulement si les jalons sont atteints : ≥ 25 domaines payants, 2 logisticiens intégrés, 1 lettre d'intention de licence.", 14)

# 3 — Why
s = d.slide("Start with why : aucun vin ne devrait rester à quai", "Vision & motivation")
d.text(s, 0.6, 1.55, 6.3, 3.8, ["Au XIXᵉ siècle, aucun vin ne circulait sans « passavant », le laissez-passer des douanes.", "",
                               "Aujourd'hui, un domaine de Loire, de Vénétie ou de la Rioja qui vend en ligne à un client suisse, belge ou allemand perd encore du temps — et parfois la vente — entre sa boutique et son logisticien.", "",
                               "Passavin veut que chaque bouteille vendue en ligne en Europe puisse partir, sans ressaisie, vers toutes les destinations réellement livrables."], 16)
d.box(s, 7.3, 1.55, 5.4, 5.1, WHITE)
d.text(s, 7.6, 1.7, 5, 0.4, "SOURCES DE MOTIVATION ET APPROCHE", 12, TEAL, True)
d.bullets(s, 7.6, 2.15, 4.9, 4.4, ["Expérience terrain du porteur : commercial dans le transport de vin pour particuliers, partout dans le monde",
                                   "Constat répété d'intégrations refaites « une par une » chez les logisticiens (signal observé, pas des clients de Passavin)",
                                   "Ambition : une PME SaaS européenne, neutre entre logisticiens",
                                   "Pragmatisme : un CMS à la fois, un pays à la fois, chaque recrutement déclenché par un jalon",
                                   "Tester la demande avant de coder : seuils GO / STOP fixés à l'avance"], 14)

# 4 — Environnement
s = d.slide("Un environnement qui pousse les domaines vers la vente directe en Europe", "Environnement & pourquoi maintenant")
d.kpis(s, 0.6, 1.5, [("−8 %", "exportations françaises de vins & spiritueux en 2025 (14,3 Md€, FEVS)"), ("−3,7 %", "exportations italiennes de vin 2025 (7,77 Md€, ISTAT)"),
                     ("−5 %", "exportations espagnoles de vin 2025 (2,88 Md€, AEAT/OIVE)"), ("46 %", "du CA des domaines allemands en vente directe (DWI/Geisenheim 2024)")], w=2.85, lab=11)
d.table(s, 0.6, 3.3, 12.1, [["Tendance", "Donnée [FAIT]", "Effet pour Passavin"],
                            ["Droits américains de 15 % depuis août 2025", "États-Unis : 1ᵉʳ marché du vin français (20 %)", "Report vers l'Europe, où l'expédition aux particuliers est possible"],
                            ["Vente à distance intra-UE depuis 2023", "Accises du pays de destination (directive UE 2020/262)", "Besoin de logisticiens spécialisés → besoin de connecteurs"],
                            ["Digitalisation des producteurs", "79 % des viticulteurs FR ont un site, 21 % vendent en ligne (Afnic 2025)", "Base installée de boutiques à connecter"],
                            ["Les logisticiens se lancent seuls", "MBE Wine for Europe (juin 2026), plugin Hillebrand Gori", "Fenêtre ouverte pour un acteur neutre, mais pas pour longtemps"]],
        widths=[3.4, 4.5, 4.2], size=12, h_row=0.55)
d.text(s, 0.6, 6.15, 12.1, 0.7, "Point d'attention : les budgets des domaines se resserrent aussi. Prix d'entrée bas, gain démontré en démo.", 13, WINE, True)

# 5 — Problème
s = d.slide("Le problème : trois frictions entre « vendre » et « livrer »", "Problème")
for j, (t, b) in enumerate([("01 · Avant la vente", "La boutique propose des pays où aucun logisticien ne livre le vin, avec un forfait de port approximatif → commande annulée ou vendue à perte."),
                            ("02 · Après la vente", "La commande est ressaisie à la main dans le portail du logisticien : ≈ 12 min par commande étrangère [hypothèse à mesurer]."),
                            ("03 · Après l'expédition", "Le numéro de suivi est recopié à la main ; le client inquiet écrit ; les incidents se gèrent par mail.")]):
    d.box(s, 0.6 + j * 4.1, 1.6, 3.85, 3.3, WHITE)
    d.text(s, 0.85 + j * 4.1, 1.8, 3.4, 0.5, t, 18, WINE, True)
    d.text(s, 0.85 + j * 4.1, 2.4, 3.4, 2.4, b, 15)
d.callout(s, 0.6, 5.2, 12.1, 1.4, "Deux victimes : le domaine (temps, ventes perdues) et le logisticien, qui paie chaque intégration sur mesure — ou un plugin maison peu adopté. D'où deux offres : Direct pour les domaines, Partner pour les logisticiens.", 15)

# 6 — Prouvé / pas prouvé
s = d.slide("Le problème est réel, son ampleur n'est pas encore prouvée", "Problème — validé / non validé")
d.table(s, 0.6, 1.5, 12.1, [["Affirmation", "Statut", "Comment le prouver"],
                            ["Les intégrations logistiques sont faites au cas par cas", "Observé chez un logisticien (signal, pas des clients)", "Entretiens avec 4 logisticiens"],
                            ["Les outils génériques gèrent mal l'alcool international", "Vérifié (Packlink, FedEx)", "—"],
                            ["Un plugin vin existe mais n'est pas adopté", "Vérifié (Hillebrand Gori : < 10 installations)", "Comprendre pourquoi (3 entretiens)"],
                            ["≈ 12 min de ressaisie par commande étrangère", "**Hypothèse**", "Chronométrage chez 25 domaines"],
                            ["25 % des caves italiennes et 50 % des bodegas vendent sur leur site", "**Hypothèse (Nomisma ≈ 50 %, EAE 70 %)", "10 entretiens IT, 8 ES, scan de sites"],
                            ["Les domaines paieront 94 €/mois en moyenne", "**Hypothèse**", "Van Westendorp + 8 pré-ventes avec acompte"],
                            ["Un logisticien paiera une licence de marque blanche", "**Hypothèse**", "1 lettre d'intention avant la seed"]],
        widths=[5.0, 3.9, 3.2], size=12, h_row=0.6)

# 7 — Solution
s = d.slide("La solution : votre boutique, branchée à vos logisticiens vin", "Solution")
d.callout(s, 0.6, 1.5, 12.1, 0.95, "« Au panier, vos clients étrangers ne voient que les destinations réellement livrables, au bon prix. Après l'achat, la commande part chez le logisticien et le suivi revient tout seul. »", 15)
d.flow(s, 0.6, 2.75, [("Installation", "Connexion du logisticien, contrôle des fiches produits"), ("Checkout", "Éligibilité pays + tarif via les API des logisticiens"),
                      ("Paiement", "Passerelle du domaine (aucun flux financier chez nous)"), ("Pré-contrôle", "Adresse, téléphone, colis, valeur ; file d'anomalies"),
                      ("Transmission", "Ordre d'expédition chez le logisticien choisi"), ("Suivi", "Statuts récupérés, mail au client"), ("Pilotage", "Délais, coûts, incidents par pays")], w=1.6, h=1.6, size=11)
d.bullets(s, 0.6, 4.75, 12.1, 2.0, ["**Ne promet pas une conformité automatique : les règles d'éligibilité sont fournies et garanties par chaque logisticien ; Passavin les applique.",
                                    "Pas d'envoi direct à des particuliers aux États-Unis (système three-tier).",
                                    "RGPD : Passavin est sous-traitant du domaine ; hébergement UE ; seules les données utiles à l'expédition sont transmises au logisticien choisi."], 14)

# 8 — Produits
s = d.slide("Deux offres, un seul moteur, trois CMS", "Produits et services")
d.card(s, 0.6, 1.5, 5.9, 2.75, "PASSAVIN DIRECT — pour les domaines", ["Extension CMS + moteur cloud multi-logisticiens", "Essentiel 49 € · Pro 99 € · Expert 199 €/mois", "Installation 490 € (0 € via une agence)", "Domaine libre de choisir et de changer de logisticien"], tcolor=WINE, size=13)
d.card(s, 6.8, 1.5, 5.9, 2.75, "PASSAVIN PARTNER — pour les logisticiens", ["Même moteur, à la marque du logisticien", "Licence 690 / 1 490 / 2 490 €/mois + mise en place", "0,40 à 0,60 € par expédition", "Le logisticien l'offre à ses domaines clients"], tcolor=TEAL, size=13)
d.text(s, 0.6, 4.45, 12, 0.4, "FEUILLE DE ROUTE PRODUIT", 12, TEAL, True)
d.flow(s, 0.6, 4.85, [("M1-M4", "MVP WooCommerce, 2 adaptateurs logisticiens"), ("M5", "Commercialisation France"), ("M15", "1ʳᵉ licence Partner"), (f"M{P['cms_launch']['shopify']}", "App Shopify (0 % de commission jusqu'à 1 M$)"), (f"M{P['cms_launch']['presta']}", "Module PrestaShop (vendu en direct)"), ("M26+", "Nouveaux adaptateurs, marque blanche Premium")], w=1.9, h=1.75, size=11)

# 9 — Marché
s = d.slide(f"Un marché européen d'environ {n(round(SAM_TOT, -2))} domaines adressables", "Marché (market size)")
rows = [["Pays", "Producteurs (base)", "Vendent en ligne", "TAM", "SAM (3 CMS × export)", "SOM M36", "Pénétration"]]
base = {"FR": "43 000 [FAIT]", "IT": "17 000 [FAIT ICQRF]", "ES": "3 925 [FAIT INE]", "DE": "15 151 [FAIT Destatis]", "PT": "3 000 [HYP, NV]"}
onl = {"FR": "21 % [FAIT]", "IT": "25 % [HYP]", "ES": "50 % [HYP]", "DE": "25 % [HYP]", "PT": "25 % [HYP]"}
lab_c = {"FR": "France", "IT": "Italie", "ES": "Espagne", "DE": "Allemagne", "PT": "Portugal"}
for c in ("FR", "IT", "ES", "DE", "PT"):
    rows.append([lab_c[c], base[c], onl[c], n(MK[c]["tam"]), n(MK[c]["sam_3cms"]), n(CNT[c]), f"{CNT[c] / MK[c]['sam_3cms']:.0%}".replace("%", " %")])
rows.append(["**Total", "", "", f"**{n(TAM_TOT)}", f"**{n(SAM_TOT)}", f"**{n(Y[2]['clients_fin'])}", f"**{Y[2]['clients_fin'] / SAM_TOT:.0%}".replace("%", " %")])
d.table(s, 0.6, 1.5, 12.1, rows, widths=[1.6, 2.4, 1.9, 1.3, 2.2, 1.3, 1.4], size=12, h_row=0.45)
d.bullets(s, 0.6, 4.85, 12.1, 2.0, [f"En valeur (≈ 1 300 €/an par domaine) : TAM ≈ {TAM_TOT * 1300 / 1e6:.0f} M€/an, SAM ≈ {SAM_TOT * 1300 / 1e6:.1f} M€/an".replace(".", ",") + f" ; + {n(Y[2]['wl_domaines'])} domaines activés par les logisticiens revendeurs.",
                                    "Parts de CMS : WMTips 2026 (tous secteurs, faute de données propres au vin). Parts « export particuliers » : hypothèses à mesurer.",
                                    "**Allemagne = surtout un pays de destination : pas de commercial dédié. Espagne = vivier le plus petit : poste partagé avec le Portugal."], 13)

# 10 — Personas
s = d.slide("Cibles : quatre personas (illustratifs)", "Clientèle — buyer personas")
for j, (t, who, pain, msg) in enumerate([
        ("Claire, 42 ans", "Vigneronne-gérante (Loire), WooCommerce", "Ressaisies le soir ; « où est mon colis ? »", "« Vos commandes étrangères partent sans ressaisie. »"),
        ("Marco, 38 ans", "Responsable commercial, Valpolicella", "Pays non livrables au panier, accises mal estimées", "« Plus aucun pays impossible au panier. »"),
        ("Sophie, 38 ans", "Dirigeante d'agence web, 25 domaines", "Intégrations refaites à chaque client", "« Une intégration maintenue, 20-25 % récurrents. »"),
        ("Marc, 50 ans", "Directeur commercial d'un logisticien", "Plugin maison coûteux, peu adopté", "« Votre connecteur à votre marque en 8 semaines. »")]):
    x = 0.6 + j * 3.05
    d.box(s, x, 1.5, 2.85, 4.6, WHITE)
    d.text(s, x + 0.2, 1.65, 2.5, 0.45, t, 17, WINE, True)
    d.text(s, x + 0.2, 2.1, 2.5, 0.8, who, 12, GREY)
    d.text(s, x + 0.2, 2.95, 2.5, 0.35, "FRUSTRATION", 10, TEAL, True); d.text(s, x + 0.2, 3.25, 2.5, 1.0, pain, 13)
    d.text(s, x + 0.2, 4.3, 2.5, 0.35, "MESSAGE", 10, TEAL, True); d.text(s, x + 0.2, 4.6, 2.5, 1.4, msg, 13, INK, True)
d.text(s, 0.6, 6.3, 12.1, 0.5, "ICP Direct : domaine de 10-60 ha, boutique active depuis ≥ 1 an, 10-80 commandes/mois dont ≥ 20 % hors du pays. Personas fictifs, à remplacer par les entretiens.", 12, GREY)

# 11 — Concurrence
s = d.slide("Concurrence : l'idée existe… en partie, et personne n'est neutre", "Analyse de la concurrence")
d.table(s, 0.6, 1.45, 12.1, [["Acteur", "Type", "Force", "Limite pour notre cible"],
                             ["Hillebrand Gori eShipping (DHL)", "Logisticien + plugin", "Droits au checkout, FR/ES/IT", "Mono-logisticien, < 10 installations Woo, pas d'app Shopify trouvée"],
                             ["MBE Wine for Europe (juin 2026)", "Logisticien + module", "Accises UE, réseau Mail Boxes Etc.", "Mono-réseau → aussi un client Partner potentiel"],
                             ["Tannico WinePlatform (IT)", "E-shop en marque blanche + logistique", "Clé en main pour les caves italiennes", "Remplace la boutique du domaine"],
                             ["VinTrail, Baqio, Wineriz", "Plateformes de vente vin", "Tout-en-un", "Remplacent la boutique ; modèle fermé"],
                             ["Sendcloud, Boxtal, Packlink, ShipStation", "Multi-transporteurs génériques", "Matures, pas chers", "Alcool international mal couvert (Packlink : pas de vin hors UE)"],
                             ["Commerce7, Vinoshipper, Sovos", "Conformité DTC", "Référence aux États-Unis", "Marché américain, pas l'export depuis l'UE"],
                             ["Marketplaces (Vivino, Plugwine…)", "Vente intermédiée", "Audience", "Le domaine perd la relation client (≈ 20 % de commission)"],
                             ["Développement sur mesure / agences", "Substitut", "Sur mesure", "Non maintenu, refait à chaque client"]],
        widths=[3.0, 2.6, 2.8, 3.7], size=11, h_row=0.52)
d.text(s, 0.6, 6.25, 12.1, 0.5, "Sources : sites des éditeurs, wordpress.org, Vinetur (2026). Acteurs non retrouvés lors de la recherche : non cités.", 11, GREY)

# 12 — USP
s = d.slide("Positionnement et proposition de valeur", "Avantages concurrentiels — USP")
d.callout(s, 0.6, 1.5, 12.1, 1.1, "USP : « Le connecteur neutre qui branche votre boutique à plusieurs logisticiens vin. » Jamais « le seul plugin vin » : Hillebrand Gori en a un, mono-logisticien.", 16)
for j, (t, b) in enumerate([("Neutre", "Plusieurs logisticiens, aucun classement payant, mêmes conditions pour tous (Planet Aura compris)"),
                            ("Spécialisé alcool", "Éligibilité des destinations et tarifs réels au checkout, là où les outils génériques échouent"),
                            ("Multi-CMS", "WooCommerce, Shopify, PrestaShop : 100 % du vivier couvert à M25"),
                            ("Deux faces", "Domaines (Direct) et logisticiens (Partner) : chaque logisticien devient un canal")]):
    d.card(s, 0.6 + j * 3.05, 2.9, 2.85, 2.3, t, b, size=13)
d.bullets(s, 0.6, 5.45, 12.1, 1.4, ["Barrières à construire : réseau de logisticiens intégrés, base de règles d'éligibilité maintenue, réseau de 27 agences, neutralité contractuelle.",
                                    "Barrière faible au départ : la vitesse d'exécution et la connaissance du terrain du fondateur. À dire franchement."], 13)

# 13 — Business model
s = d.slide("Business model : deux offres, sept lignes de revenus", "Business model")
d.table(s, 0.6, 1.45, 12.1, [["Offre", "Ligne de revenus", "Qui paie", "Prix", "CA an 3 (central)"],
                             ["Direct", "Abonnement", "Domaine", "49 / 99 / 199 €/mois (ARPA 94 €)", k(Y[2]["ca_abonnements"])],
                             ["Direct", "Installation accompagnée", "Domaine", "490 € (190 € pilotes, 0 € via agence)", k(Y[2]["ca_installation"])],
                             ["Direct", "Frais de plateforme (à valider)", "Logisticien choisi", "1,50 €/expédition", k(Y[2]["ca_logisticiens"])],
                             ["Direct", "Services (≤ 15 % du CA)", "Domaine", "Sur devis", k(Y[2]["ca_services"])],
                             ["Partner", "Licence de marque blanche", "Logisticien", "690 / 1 490 / 2 490 €/mois", k(Y[2]["ca_wl_licences"])],
                             ["Partner", "Mise en place", "Logisticien", "4 000 à 12 000 €", k(Y[2]["ca_wl_setup"])],
                             ["Partner", "Frais par expédition", "Logisticien", "0,40 à 0,60 €", k(Y[2]["ca_wl_expeditions"])],
                             ["**Total", "", "", "", f"**{k(Y[2]['ca_total'])}"]],
        widths=[1.3, 3.3, 2.1, 3.6, 1.8], size=12, h_row=0.47)
d.text(s, 0.6, 5.85, 12.1, 0.9, "Abandonnés : création de boutiques (confiée aux agences, devenues un canal) et développements non réutilisables. Abonnement annuel = 2 mois offerts (lisse la saisonnalité).", 13, GREY)

# 14 — Économie unitaire + prix Partner
s = d.slide("Économie unitaire et prix de la marque blanche", "Business model — tarification")
d.kpis(s, 0.6, 1.5, [(f"{U['contribution']:.0f} €", "contribution mensuelle par domaine (abonnement + frais − coûts variables)"),
                     (f"{U['ltv']:,.0f} €".replace(",", " "), f"LTV (durée de vie {U['duree_vie']:.0f} mois, résiliation 1,5 %/mois)"),
                     (f"{U['cac_a3']:,.0f} €".replace(",", " "), f"CAC an 3 ({n(Y[0]['cac'])} € en an 1, {n(Y[1]['cac'])} € en an 2)"),
                     (f"{U['ltv_cac']:.1f}".replace(".", ","), f"LTV / CAC ; CAC récupéré en {U['payback']:.0f} mois")], w=2.85, lab=11)
d.text(s, 0.6, 3.3, 6, 0.4, "POURQUOI UN LOGISTICIEN PAIERAIT 1 490 €/MOIS", 12, TEAL, True)
d.table(s, 0.6, 3.7, 6.6, [["Coût pour 2 CMS", "1ʳᵉ année", "Ensuite / an"],
                           ["Plugin développé en interne [HYP]", "≈ 69 k€", "≈ 15 k€"],
                           ["Licence Business + mise en place", "25,9 k€", "17,9 k€"],
                           ["Risque d'adoption", "Le sien", "Partagé"]], widths=[3.2, 1.7, 1.7], size=12, h_row=0.5)
d.bullets(s, 7.5, 3.4, 5.2, 3.4, ["Repères publics 2026 : Sendcloud Pro 639 €/mois + 0,07 €/étiquette ; Cargoson 199 à 1 000 $/mois.",
                                  "Passavin Partner est au-dessus d'un outil d'étiquetage : éligibilité alcool, 3 CMS, marque du logisticien.",
                                  "0,50 €/expédition en marque blanche contre 1,50 € en mode neutre : en marque blanche, c'est le logisticien qui apporte les clients."], 13)

# 15 — Étapes clés
s = d.slide("Stratégie : un pays et un CMS à la fois", "Stratégie — étapes clés")
d.flow(s, 0.6, 1.55, [("Nov. 26-fév. 27", "Validation FR : 25 entretiens, 8 pré-ventes, 2 logisticiens"), ("M1-M12", f"MVP Woo, pilotes, France : {n(R[11]['clients'])} domaines ; seed M12"),
                      ("M13-M18", f"Partner manager, Italie (M{H['sales_it']}), Shopify (M18) : {n(R[17]['clients'])} domaines"),
                      ("M19-M24", f"Espagne-Portugal (M{H['sales_es']}) : {n(R[23]['clients'])} domaines, 2 licences"),
                      ("M25-M30", f"PrestaShop, customer success IT/ES : {n(R[29]['clients'])} domaines"),
                      ("M31-M36", f"Allemagne via le réseau : {n(R[35]['clients'])} domaines + {n(R[35]['wl_domains'])} en marque blanche")], w=1.95, h=2.1, size=11)
d.image(s, CH("pays.png"), 0.6, 3.8, h=3.05)
d.bullets(s, 6.9, 3.95, 5.8, 2.9, ["Chaque étape a un jalon mesurable : si le jalon n'est pas atteint, l'étape suivante attend.",
                                   "Italie avant l'Espagne : SAM 2,3 fois plus grand.",
                                   "PrestaShop avancé à M25 : il pèse plus que Shopify en France (27 % contre 13 %) et en Espagne (20 % contre 6 %)."], 13)

# 16 — Stratégie commerciale
s = d.slide("Stratégie commerciale : canaux, cycle de vente, objectifs", "Stratégie commerciale — chiffres clés")
d.table(s, 0.6, 1.45, 7.3, [["Objectif (central)", "2027", "2028", "2029"],
                            ["Nouveaux domaines Direct", n(Y[0]["nouveaux"]), n(Y[1]["nouveaux"]), n(Y[2]["nouveaux"])],
                            ["Domaines actifs fin d'année", n(Y[0]["clients_fin"]), n(Y[1]["clients_fin"]), n(Y[2]["clients_fin"])],
                            ["Licences / domaines marque blanche", "0 / 0", f"{Y[1]['wl_partenaires']} / {n(Y[1]['wl_domaines'])}", f"{Y[2]['wl_partenaires']} / {n(Y[2]['wl_domaines'])}"],
                            ["Revenu récurrent annualisé", k(Y[0]["arr_fin"]), k(Y[1]["arr_fin"]), k(Y[2]["arr_fin"])],
                            ["Agences / apporteurs actifs", f"{Y[0]['agences_fin']} / {Y[0]['apporteurs_fin']}", f"{Y[1]['agences_fin']} / {Y[1]['apporteurs_fin']}", f"{Y[2]['agences_fin']} / {Y[2]['apporteurs_fin']}"],
                            ["Effectif", n(Y[0]["effectif_fin"]), n(Y[1]["effectif_fin"]), n(Y[2]["effectif_fin"])]],
        widths=[3.4, 1.3, 1.3, 1.3], size=12, h_row=0.48)
d.text(s, 8.2, 1.45, 4.5, 0.4, "CYCLE DE VENTE (DOMAINE)", 12, TEAL, True)
d.table(s, 8.2, 1.85, 4.5, [["Étape", "Taux"], ["Prospect → contact", "30 %"], ["Contact → démo", "50 %"], ["Démo → essai", "50 %"], ["Essai → client", "50 %"], ["**100 prospects → 3,75 clients", "**6-8 sem."]],
        widths=[3.2, 1.3], size=12, h_row=0.46)
d.bullets(s, 0.6, 5.05, 12.1, 1.8, ["Un·e sales à pleine productivité fait ≈ 30 démos et contacte ≈ 200 prospects par mois → un·e alternant·e SDR prend les rendez-vous.",
                                    "Logisticiens : cycle de 6 à 9 mois (comité de direction, essai technique), porté par le dirigeant et le partner manager.",
                                    "Saisonnalité : prospection de janvier à juin ; vendanges = agences et contenus ; Noël = support."], 13)

# 17 — Marketing
s = d.slide("Stratégie marketing : budget, événements, contenus", "Stratégie marketing")
d.table(s, 0.6, 1.45, 6.2, [["k€", "An 1", "An 2", "An 3"], ["Marketing (salons, contenus, tests)", "12", "40", "60"], ["Programme partenaires", "1,5", "8", "15"],
                            ["Déplacements commerciaux", "3", "10", "18"], ["Localisation IT/ES/DE/PT", "0", "12", "8"], ["**Total", "**16,5", "**70", "**101"]],
        widths=[3.2, 1.0, 1.0, 1.0], size=12, h_row=0.5)
d.table(s, 7.1, 1.45, 5.6, [["Événement", "Usage"], ["Vinitech, Wine Paris", "Validation, pré-ventes, logisticiens"], ["ProWein (Düsseldorf)", "Logisticiens européens"],
                            ["Vinitaly 2028 (Vérone)", "Lancement Italie"], ["Fenavin / Barcelona Wine Week", "Lancement Espagne"], ["AG des syndicats d'appellation", "Ateliers de 20 min"]],
        widths=[2.6, 3.0], size=12, h_row=0.5)
d.bullets(s, 0.6, 4.75, 12.1, 2.0, ["Inbound : fiche wordpress.org (version « suivi seul » gratuite), App Store Shopify, guide « Expédier du vin aux particuliers en Europe depuis 2023 », webinaires co-animés avec des logisticiens.",
                                    "Pas de stand payant avant l'an 3 : la valeur est dans les rendez-vous préparés. Dates des salons à reconfirmer.",
                                    "Pourquoi maintenant : exportations en baisse, accises intra-UE depuis 2023, logisticiens qui lancent leurs propres modules (fenêtre pour un acteur neutre)."], 13)

# 18 — Équipe fondatrice
s = d.slide("L'équipe : un fondateur commercial, une équipe à construire par jalons", "Équipe — rôles et complémentarités")
d.table(s, 0.6, 1.45, 12.1, [["Rôle", "Qui", "Apport", "Complémentarité"],
                             ["Dirigeant·e — vente, partenariats, management", "[Nom Prénom], commercial dans le transport de vin", "Terrain, logisticiens, domaines", "Porte la vision et les licences Partner"],
                             ["Lead développeur·se / CTO (M10)", "**À recruter — priorité n°1 (BSPCE)", "Architecture SaaS, CMS, API", "Compense l'absence de compétence technique"],
                             ["3 sales pays (M10, M17, M20)", "À recruter (natif·ves IT, ES)", "Langue, réseau local", "Le dirigeant passe de vendeur à manager"],
                             ["Partner manager (M13)", "À recruter", "Animation de réseaux", "Démultiplie la vente (agences, apporteurs)"],
                             ["Customer success ×2 (M16, M26)", "À recruter", "Activation, rétention", "Protège le revenu récurrent"],
                             ["Conseils externes", "Avocat, expert-comptable, CPI", "Contrats, PI, finance", "Sécurise sans coût fixe"]],
        widths=[3.3, 3.1, 2.6, 3.1], size=12, h_row=0.6)
d.text(s, 0.6, 5.9, 12.1, 0.9, "Aucun associé, aucun diplôme, aucune compétence n'est inventé. Formation du fondateur prévue : finance d'une SAS, management d'une première équipe, contrats SaaS et RGPD.", 13, GREY)

# 19 — Organigramme + plan de recrutement
s = d.slide("Organigramme sur 3 ans et plan de recrutement relié au BP", "Équipe — organisation")
for j, (t, body, eff) in enumerate([
        ("Fin 2027", ["Dirigeant·e", f"Lead dev (M{H['lead_dev']})", f"Sales France (M{H['sales_fr']})", "Réseau : 2 agences, 3 apporteurs", "Freelance, avocat, comptable"], Y[0]["effectif_fin"]),
        ("Fin 2028", ["Partner manager, alternant·e, dév. Shopify (M13)", f"Customer success FR/EN (M{H['csm1']})", f"Sales Italie (M{H['sales_it']}), Sales ES-PT (M{H['sales_es']})", f"Réseau : {Y[1]['agences_fin']} agences, {Y[1]['apporteurs_fin']} apporteurs, {Y[1]['wl_partenaires']} licences"], Y[1]["effectif_fin"]),
        ("Fin 2029", [f"Customer success IT/ES (M{H['csm2']})", f"Dév. marque blanche (M{H['dev3']})", "4 pôles : produit, ventes, réseau, clients", f"Réseau : {Y[2]['agences_fin']} agences, {Y[2]['apporteurs_fin']} apporteurs, {Y[2]['wl_partenaires']} licences"], Y[2]["effectif_fin"])]):
    x = 0.6 + j * 4.1
    d.box(s, x, 1.45, 3.85, 2.75, WHITE)
    d.text(s, x + 0.2, 1.55, 2.5, 0.45, t, 17, WINE, True); d.text(s, x + 2.4, 1.55, 1.3, 0.45, f"{eff} pers.", 15, TEAL, True, align=PP_ALIGN.RIGHT)
    d.bullets(s, x + 0.2, 2.05, 3.5, 2.1, body, 12, gap=2)
d.table(s, 0.6, 4.4, 12.1, [["Poste", "Jalon (domaines actifs)", "Arrivée centrale / pessimiste", "Ligne du BP produite"],
                            ["Sales Italie", "60", f"M{H['sales_it']} / M{PE['hired']['sales_it']}", f"{n(tot('sales_it'))} domaines"],
                            ["Sales Espagne & Portugal", "90", f"M{H['sales_es']} / M{PE['hired']['sales_es']}", f"{n(tot('sales_es'))} domaines"],
                            ["Partner manager", "25", f"M{H['pm']} / M{PE['hired']['pm']}", f"{n(AG)} via agences, {n(tot('apporteur'))} via apporteurs, 4 licences"],
                            ["Customer success IT/ES + dév. marque blanche", "200", f"M{H['csm2']} / non recrutés", "Rétention IT/ES, licences 3 et 4"]],
        widths=[3.6, 2.3, 2.8, 3.4], size=11, h_row=0.42)
d.text(s, 0.6, 6.55, 12.1, 0.35, "Règle de gestion codée dans le modèle : pas d'embauche si la trésorerie couvre moins de 6 mois de burn.", 11, GREY)

# 20 — Équipe commerciale interne
s = d.slide("L'équipe commerciale : rôles, objectifs SMART, rémunération", "Management — équipe interne")
d.table(s, 0.6, 1.45, 12.1, [["Poste", "Rôle", "Objectifs SMART (central)", "Rémunération"],
                             ["Dirigeant·e", "Vente an 1, licences, management", f"{n(src(0,'ceo'))} domaines en 2027 ; seed à M12 ; 1ʳᵉ licence à M15", "0 € an 1 ; 2 500 puis 3 500 € brut/mois"],
                             ["Sales France", "Chasse + comptes jusqu'au customer success", f"{n(src(1,'sales_fr'))} domaines en 2028, {n(src(2,'sales_fr'))} en 2029 ; ≥ 30 démos/mois ; ≥ 25 % de conversion", "32 k€ + 15 % de la 1ʳᵉ année signée (≈ 170 €/domaine)"],
                             ["Sales Italie", "Ouvrir l'Italie, recruter les agences", f"{n(src(1,'sales_it'))} domaines (M{H['sales_it']}-M24), {n(src(2,'sales_it'))} en 2029 ; 5 agences à M34", "Idem, via Employer of Record (+450 €/mois)"],
                             ["Sales Espagne & Portugal", "Ouvrir l'Espagne puis le Portugal", f"{n(src(1,'sales_es'))} domaines (M{H['sales_es']}-M24), {n(src(2,'sales_es'))} en 2029", "Idem"],
                             ["Partner manager", "Agences, apporteurs, licences", f"{Y[1]['agences_fin']} → {Y[2]['agences_fin']} agences ; {n(AG)} domaines via agences sur 3 ans ; 4 licences", "38 k€ + 1 500 €/licence + 50 €/domaine agence"],
                             ["Customer success ×2", "Activation, support, montée en gamme", "Résiliation ≤ 1,5 %/mois ; activation ≥ 80 % ; ARPA ≥ 94 €", "32 k€ dont prime de rétention ≤ 2 k€"],
                             ["Alternant·e SDR", "Rendez-vous, contenus", "25 démos qualifiées/mois", "≈ 1 000 €/mois net d'aides"]],
        widths=[2.0, 3.0, 4.4, 2.7], size=11, h_row=0.6)
d.text(s, 0.6, 6.4, 12.1, 0.45, "Variable payé chaque trimestre, récupéré si le client résilie avant 3 mois, plafonné à 200 % ; BSPCE (réserve de 10 %) pour les premiers salariés.", 12, GREY)

# 21 — Réseau externe
s = d.slide("Le réseau commercial externe : agences, apporteurs, logisticiens", "Management — réseau")
d.table(s, 0.6, 1.45, 12.1, [["", "Agences Silver", "Agences Gold", "Apporteurs d'affaires", "Logisticiens revendeurs"],
                             ["Rôle", "Recommandent et installent", "Idem + co-marketing", "Consultants, œnologues, syndicats", "Offrent Passavin Partner à leurs domaines"],
                             ["Objectif SMART", "≥ 1 domaine/trimestre", "≥ 1,2 domaine/mois", "3 domaines/an", "≥ 40 domaines activés la 1ʳᵉ année"],
                             ["Rémunération", "20 % récurrent à vie", "25 % récurrent + 500 €/an", "250 €/domaine actif à 3 mois", "Ils paient la licence ; dégressivité"],
                             ["Recrutement", "Annuaires Shopify/PrestaShop, sites de domaines", "Promotion des Silver", "Réseau, AG de syndicats", "Liste de 15 logisticiens européens"],
                             ["Formation", "Certification 2 h (FR/IT/ES)", "+ formation avancée", "Kit + 45 min", "2 sessions pour leurs commerciaux"],
                             ["Animation", "Newsletter, webinaire, CRM", "+ revue trimestrielle", "Réunion semestrielle", "Comité de pilotage trimestriel"],
                             ["Sous-performance", "0 en 6 mois → plan 90 j, puis inactive", "2 trimestres sous seuil → Silver", "6 mois sans lead → en sommeil", "Plan conjoint ; licence minimale due"],
                             ["Fin 2029 (central)", f"{R[-1]['n_silver']} actives · {n(tot('agence_silver'))} domaines", f"{R[-1]['n_gold']} actives · {n(tot('agence_gold'))} domaines", f"{Y[2]['apporteurs_fin']} actifs · {n(tot('apporteur'))} domaines", f"{Y[2]['wl_partenaires']} licences · {n(Y[2]['wl_domaines'])} domaines"]],
        widths=[1.8, 2.6, 2.4, 2.6, 2.7], size=11, h_row=0.56)
d.text(s, 0.6, 6.55, 12.1, 0.35, "Syndicats et associations : pas de commission individuelle mais un tarif adhérent (−10 %), pour éviter tout conflit d'intérêts.", 11, GREY)

# 22 — Recrutement, intégration, animation
s = d.slide("Recruter, intégrer, former, animer", "Management — processus")
d.card(s, 0.6, 1.45, 3.9, 2.55, "Recrutement", ["Poste ouvert seulement si le jalon est atteint", "WTTJ, LinkedIn, InfoJobs IT/ES, cooptation", "Jeu de rôle : démo à un vigneron sceptique", "Prise de références"], size=12)
d.card(s, 4.7, 1.45, 3.9, 2.55, "Intégration 30-60-90 jours", ["Semaine 1 : immersion chez un domaine et un logisticien", "J+30 : démo certifiée", "J+60 : quota à 66 %", "J+90 : quota plein (rampe 3-4 mois)"], size=12)
d.card(s, 8.8, 1.45, 3.9, 2.55, "Formation continue", ["Playbook commercial (Notion)", "Coaching sur appels enregistrés", "1 000 €/an/salarié", "Mise à jour réglementaire semestrielle"], size=12)
d.table(s, 0.6, 4.2, 7.6, [["Rituel", "Fréquence", "Objet"], ["Revue de pipeline", "Hebdo, 45 min", "Prévision, blocages"], ["1:1", "Hebdo (sales)", "Coaching, moral"],
                           ["Revue d'activité", "Mensuelle", "Tableau de bord"], ["OKR", "Trimestriels", "Objectifs ≤ 3/pers."], ["Revue partenaires", "Trimestrielle", "Agences Gold, logisticiens"]],
        widths=[2.4, 2.0, 3.2], size=11, h_row=0.42)
d.card(s, 8.5, 4.2, 4.2, 2.55, "Outils et challenges", ["CRM HubSpot : pipelines par pays, portail partenaires", "Metabase + Stripe : tableau de bord", "Challenge « Pro d'abord » trimestriel", "« Partenaire de l'année » (voyage salon)"], size=12)

# 23 — Style de management
s = d.slide("Style de management, motivation, sous-performance", "Management — leadership")
d.table(s, 0.6, 1.45, 6.6, [["Style (Hersey-Blanchard)", "Pour qui, quand"], ["S1 Directif", "Alternant·e ; sales en rampe (0-3 mois)"], ["S2 Persuasif", "Sales IT et ES, 6 premiers mois"],
                            ["S3 Participatif", "Customer success, développeurs"], ["S4 Délégatif", "Sales FR après 12 mois ; partner manager"]],
        widths=[2.6, 4.0], size=12, h_row=0.48)
d.text(s, 0.6, 4.0, 6.6, 0.4, "OKR trimestriels reliés au BP — exemple T1 2029", 12, TEAL, True)
d.text(s, 0.6, 4.4, 6.6, 1.3, ["Entreprise : « Ouvrir l'Europe du Sud » → 36 domaines IT/ES, 2 agences hors de France, PrestaShop à M25.",
                               "Sales Italie : « Référence en Vénétie » → 17 domaines, 50 démos, 1 agence."], 12)
d.card(s, 7.5, 1.45, 5.2, 2.3, "Motivation et fidélisation", ["Variable juste et simple ; BSPCE", "Évolution : sales → responsable pays", "Reconnaissance, séminaire dans un domaine", "Réseau : commission récurrente à vie, Gold"], size=12)
d.card(s, 7.5, 3.95, 5.2, 2.8, "Sous-performance (équitable)", ["① Diagnostic : activité, conversion ou cause externe (CMS, saison) ?", "② Plan de progrès de 60 jours, coaching", "③ Bilan : retour à la normale, prolongation ou changement de poste", "④ Objectifs réalistes et connus : un quota manqué ne suffit pas"], size=12)
d.text(s, 0.6, 5.9, 6.6, 0.9, "Limite assumée : 7 personnes en direct pour le dirigeant à M36 → nommer un·e responsable des ventes en an 4.", 12, WINE, True)

# 24 — Démo 1
s = d.slide("Démonstration n°1 : chaque domaine vient d'une ressource identifiée", "Business plan ← objectifs commerciaux")
d.image(s, CH("sources.png"), 0.6, 1.4, w=6.4)
d.table(s, 7.2, 1.45, 5.5, [["Ressource", "Productivité", "3 ans"],
                            ["Dirigeant", "22 → 2/mois → 1/mois", n(tot("ceo"))], ["Sales France", "8/mois × couverture CMS", n(tot("sales_fr"))],
                            ["Sales Italie", "6/mois × couverture", n(tot("sales_it"))], ["Sales ES-PT", "5/mois × couverture", n(tot("sales_es"))],
                            ["Agences Silver / Gold", "0,4 / 1,2 par mois", n(AG)], ["Apporteurs", "0,25/mois", n(tot("apporteur"))],
                            ["**Total domaines Direct", "", f"**{n(NEW_TOT)}"]], widths=[2.1, 2.4, 1.0], size=11, h_row=0.42)
d.bullets(s, 0.6, 5.0, 12.1, 1.9, [f"Couverture CMS (France) : 47 % avec WooCommerce seul, 65 % avec Shopify (M18), 100 % avec PrestaShop (M25) ; rampe de 3-4 mois ; baisse quand le SAM se remplit.",
                                   f"Marque blanche : 4 licences signées par le dirigeant puis le partner manager → {n(Y[2]['wl_domaines'])} domaines activés par les logisticiens eux-mêmes.",
                                   "CA = domaines × 94 € + installations + frais logisticiens + licences : aucune ligne sans ressource."], 13)

# 25 — Compte de résultat
s = d.slide("Compte de résultat prévisionnel — scénario central", "Business plan — rentabilité")
d.table(s, 0.6, 1.45, 6.6, [["k€ HT", "2027", "2028", "2029"],
                            ["CA Direct (abonnements, installation, frais, services)", k(Y[0]["ca_direct"]), k(Y[1]["ca_direct"]), k(Y[2]["ca_direct"])],
                            ["CA Partner (marque blanche)", k(Y[0]["ca_wl"]), k(Y[1]["ca_wl"]), k(Y[2]["ca_wl"])],
                            ["**CA total", f"**{k(Y[0]['ca_total'])}", f"**{k(Y[1]['ca_total'])}", f"**{k(Y[2]['ca_total'])}"],
                            ["Marge brute", f"{Y[0]['taux_marge_brute']:.0%}".replace("%", " %"), f"{Y[1]['taux_marge_brute']:.0%}".replace("%", " %"), f"{Y[2]['taux_marge_brute']:.0%}".replace("%", " %")],
                            ["Développement externe (MVP)", k(Y[0]["developpement"]), k(Y[1]["developpement"]), k(Y[2]["developpement"])],
                            ["Salaires chargés", k(Y[0]["salaires"]), k(Y[1]["salaires"]), k(Y[2]["salaires"])],
                            ["Autres charges", k(Y[0]["autres_charges"]), k(Y[1]["autres_charges"]), k(Y[2]["autres_charges"])],
                            ["**EBITDA", f"**{k(Y[0]['ebitda'])}", f"**{k(Y[1]['ebitda'])}", f"**{k(Y[2]['ebitda'])}"],
                            ["Revenu récurrent annualisé (fin)", k(Y[0]["arr_fin"]), k(Y[1]["arr_fin"]), k(Y[2]["arr_fin"])]],
        widths=[3.3, 1.1, 1.1, 1.1], size=11, h_row=0.45)
d.image(s, CH("ca_ebitda.png"), 7.4, 1.4, w=5.3)
d.text(s, 7.4, 4.45, 5.3, 2.3, [f"Croissance : CA ×{Y[2]['ca_total'] / Y[1]['ca_total']:.1f} entre 2028 et 2029.".replace(".", ","),
                                f"EBITDA du mois 36 : {k(C['ebitda_m36'])} ; charges fixes ≈ {k(C['fixed_m36'])}/mois (11 personnes).",
                                "MVP passé en charges (prudence)."], 13)

# 26 — Scénarios & point mort
s = d.slide(f"Point mort : {BE} en central — l'objectif M36 n'est pas atteint", "Business plan — viabilité")
d.table(s, 0.6, 1.45, 6.9, [["", "Pessimiste", "Central", "Optimiste"],
                            ["Domaines Direct M36", n(PE["years"][2]["clients_fin"]), n(Y[2]["clients_fin"]), n(OP["years"][2]["clients_fin"])],
                            ["CA 2029", k(PE["years"][2]["ca_total"]), k(Y[2]["ca_total"]), k(OP["years"][2]["ca_total"])],
                            ["EBITDA du mois 36", k(PE["ebitda_m36"]), k(C["ebitda_m36"]), k(OP["ebitda_m36"])],
                            ["Point mort mensuel", "non atteint", BE, BE_OPT],
                            ["Trésorerie M36", k(PE["years"][2]["tresorerie_fin"]), k(Y[2]["tresorerie_fin"]), k(OP["years"][2]["tresorerie_fin"])]],
        widths=[2.4, 1.5, 1.5, 1.5], size=12, h_row=0.48)
d.text(s, 0.6, 4.5, 6.9, 0.4, "CE QU'IL FAUDRAIT POUR M36 (central)", 12, TEAL, True)
d.table(s, 0.6, 4.9, 6.9, [["Levier", "Central", "Requis"]] + [[l["levier"].replace("Combinaison : ", "Combinaison : "), l["central"], l["requis"]] for l in LV if l["requis"] != "insuffisant seul"][:4],
        widths=[4.3, 1.3, 1.3], size=10, h_row=0.4)
d.image(s, CH("ebitda_mensuel.png"), 7.7, 1.4, w=5.0)
d.bullets(s, 7.7, 4.1, 5.0, 2.8, [f"Sans marque blanche : point mort M{NWL['be']} ; sans frais logisticien : M{NSF['be']}.",
                                  "La résiliation seule ne suffit pas, même à 1 %.",
                                  "Pessimiste : 14 domaines à M12 → pas de seed → pivot vers la seule offre Partner, ou arrêt.",
                                  "**Je montre le vrai résultat plutôt que d'ajuster les hypothèses."], 12)

# 27 — Financement
s = d.slide(f"Besoin de financement : {k(TOTAL_FUND)} en deux temps", "Business plan — besoin financier")
d.table(s, 0.6, 1.45, 6.4, [["Source", "Montant", "Quand"], ["Apport du fondateur", k(P["equity_founder"]), "M0"], ["Prêt d'honneur", k(P["honor_loan"]), "M0-M2"],
                            ["Bourse French Tech BPI", k(P["bpi_grant"]), "M2-M4"], ["Prêt bancaire", k(P["bank_loan"]), "M2"],
                            ["**Pré-amorçage", f"**{k(PRE)}", ""], ["Seed : business angels + fonds", k(P["seed_equity"]), f"M{P['seed_month']}"],
                            ["Seed : prêt d'amorçage BPI", k(P["seed_bpi_loan"]), f"M{P['seed_month']}"], ["**Total 3 ans", f"**{k(TOTAL_FUND)}", ""]],
        widths=[3.6, 1.4, 1.4], size=12, h_row=0.44)
d.image(s, CH("tresorerie.png"), 7.3, 1.4, w=5.0)
d.bullets(s, 7.2, 4.4, 5.5, 2.5, [f"Seed = point bas sans levée + 6 mois de burn. Point bas avant seed : {k(C['cash_min'])} à M{C['cash_min_m']} → levée engagée dès M8.",
                                  "Affectation : équipe commerciale et réseau 45 %, produit 30 %, marketing Europe 10 %, customer success 5 %, juridique 3 %, réserve 7 %.",
                                  "Valorisation avant levée de 2 à 2,5 M€ [HYP] → dilution de 18 à 22 %."], 12)
d.callout(s, 0.6, 5.55, 6.4, 1.2, f"Pérennité : trésorerie positive sur 36 mois ({k(Y[2]['tresorerie_fin'])} à M36), rentable à {BE} sans nouvelle levée.", 13)

# 28 — Risques
s = d.slide("Démonstration n°2 : risques évalués et justifiés", "Risques")
d.table(s, 0.6, 1.45, 12.1, [["#", "Risque", "P", "I", "Crit.", "Prévention", "Plan B / alerte"],
                             ["R1", "Les domaines ne paient pas", "3", "4", "12", "8 pré-ventes avec acompte avant le MVP", "Pivot Partner / < 4 pré-ventes"],
                             ["R2", f"Seed non obtenue (point bas {k(C['cash_min'])} à M{C['cash_min_m']})", "3", "4", "12", "Levée dès M8, jalons publics, relais BPI", "Plan V1 France / pas de LOI à M10"],
                             ["R3", "Ouverture IT/ES plus lente", "3", "3", "9", "Recrutement par jalons, rampe de 4 mois", "Reporter l'Espagne / < 50 % du quota"],
                             ["R4", "Hypothèses IT/ES non mesurées ; Espagne saturée", "3", "3", "9", "18 entretiens IT/ES avant recrutement", "Objectifs revus / < 30 % qualifiés"],
                             ["R5", "Concurrent gratuit (MBE Wine, Hillebrand)", "3", "3", "9", "Neutralité, vitesse, agences", "En faire des clients Partner"],
                             ["R8", "Agences peu actives", "3", "3", "9", "Commission à vie, Gold, certification", "2ᵉ sales FR / < 70 % actives"],
                             ["R10", "Pas de lead dev / CTO", "3", "3", "9", "Recrutement M10 avec BSPCE", "CTO à temps partagé"],
                             ["R7", "Marque blanche vs neutralité", "2", "3", "6", "2 marques, portabilité, mêmes conditions", "Limiter au niveau Business"],
                             ["R17", "Situation du porteur vis-à-vis de son employeur", "1", "3", "3", "Projet scolaire ; aucune donnée ni code de l'employeur", "—"]],
        widths=[0.6, 3.6, 0.45, 0.45, 0.65, 3.4, 2.95], size=11, h_row=0.47)
d.text(s, 0.6, 6.25, 12.1, 0.6, "17 risques notés P × I (détail en annexe). Les 5 plus critiques sont commerciaux, financiers et d'exécution. Règle automatique : aucune embauche si la trésorerie couvre moins de 6 mois de burn.", 12, WINE, True)

# 29 — Tableau de bord
s = d.slide("Démonstration n°3 : un KPI, une cible du BP, un responsable", "Tableaux de bord")
d.table(s, 0.6, 1.45, 12.1, [["KPI", "Responsable", "2027", "2028", "2029", "Si écart"],
                             ["Nouveaux domaines Direct", "Dirigeant", n(Y[0]["nouveaux"]), n(Y[1]["nouveaux"]), n(Y[2]["nouveaux"]), "Revue canal par canal"],
                             ["Signatures Sales FR / IT / ES", "Chaque sales", n(src(0, "sales_fr")), f"{n(src(1,'sales_fr'))} / {n(src(1,'sales_it'))} / {n(src(1,'sales_es'))}", f"{n(src(2,'sales_fr'))} / {n(src(2,'sales_it'))} / {n(src(2,'sales_es'))}", "Plan de progrès 60 j"],
                             ["Domaines via agences / apporteurs", "Partner manager", f"{n(src(0,'agence_silver') + src(0,'agence_gold'))} / {n(src(0,'apporteur'))}", f"{n(src(1,'agence_silver') + src(1,'agence_gold'))} / {n(src(1,'apporteur'))}", f"{n(src(2,'agence_silver') + src(2,'agence_gold'))} / {n(src(2,'apporteur'))}", "Animation, Gold"],
                             ["Licences marque blanche", "Dirigeant + PM", "0", str(Y[1]["wl_partenaires"]), str(Y[2]["wl_partenaires"]), "Revoir l'offre Partner"],
                             ["Résiliation mensuelle", "Customer success", "≤ 1,5 %", "≤ 1,5 %", "≤ 1,5 %", "Entretien de départ"],
                             ["ARPA", "Customer success", "94 €", "94 €", "≥ 94 € (99 visé)", "Pro par défaut"],
                             ["Revenu récurrent annualisé", "Dirigeant", k(Y[0]["arr_fin"]), k(Y[1]["arr_fin"]), k(Y[2]["arr_fin"]), "Upsell, prix"],
                             ["CAC", "Dirigeant", f"{n(Y[0]['cac'])} €", f"{n(Y[1]['cac'])} €", f"{n(Y[2]['cac'])} €", "Couper le canal le plus cher"],
                             ["Runway", "Dirigeant", "≥ 6 mois", "≥ 12 mois", "≥ 6 mois", "Gel automatique des embauches"]],
        widths=[3.2, 2.0, 1.4, 1.9, 1.9, 1.7], size=11, h_row=0.47)
d.text(s, 0.6, 6.3, 12.1, 0.55, "26 KPI au total (formule, source, fréquence) dans l'annexe Excel, onglet Tableaux_de_bord : les cibles sont lues dans le même modèle que le BP.", 12, GREY)

# 30 — Annexes
s = d.slide("Annexes : validation, sources, ce qui n'est pas vérifié", "Annexes")
d.table(s, 0.6, 1.45, 5.9, [["GO seed (M10-M12)", "Seuil"], ["Domaines payants", "≥ 25"], ["Revenu récurrent mensuel", "≥ 3 000 €"], ["Résiliation", "≤ 2 %/mois"],
                            ["Logisticiens intégrés", "≥ 2"], ["Lettre d'intention de licence", "≥ 1"], ["Entretiens en Italie", "≥ 10"]],
        widths=[3.9, 2.0], size=12, h_row=0.42)
d.bullets(s, 6.8, 1.45, 5.9, 3.2, ["**Sources : FranceAgriMer, Agreste, Afnic (2025), ISTAT, ICQRF, Mediobanca, INE, OIVE/AEAT, Destatis/DWI, ViniPortugal, Eurostat, WMTips (2026), Shopify, sites des concurrents.",
                                   "Annexe financière : annexes/BP_3ans.xlsx (12 onglets : hypothèses, production par ressource, mensuel ×3, pays et CMS, marque blanche, financement, recrutement, tableaux de bord, leviers).",
                                   "Modèle reproductible : bp_model.py → build_annexes.py → build_decks.py."], 12)
d.text(s, 0.6, 4.65, 12.1, 0.4, "NON VÉRIFIÉ — À MESURER EN PHASE 1", 12, WINE, True)
d.bullets(s, 0.6, 5.05, 12.1, 1.9, ["Parts « vend en ligne » IT (25 %), ES (50 %), DE (25 %) et toutes les parts « export particuliers » ; base Portugal (3 000) ; CMS réellement utilisés par les caves.",
                                    "Disposition à payer (49-199 €), frais logisticien de 1,50 €, licence Partner ; temps de ressaisie (12 min) ; commission PrestaShop Addons.",
                                    "Chiffres lus dans la presse (ISTAT, AEAT, Mediobanca) ou dans des résumés (ISMEA) : à vérifier dans la source primaire avant usage public."], 12)
assert d.i == 30, d.i
d.save("Passavin_Dossier_Investisseur.pptx")

# =====================================================================================================
# PITCH (20 slides)
# =====================================================================================================
p = Deck("Pitch soutenance · V2")
p.cover("Vos vins passent les frontières. Vos commandes passent toutes seules.", "Pitch investisseurs — 20 minutes · Bloc 3 Euridis", "[Nom Prénom] · [Promotion] · Nom de travail provisoire")
# 2
s = p.slide("Aucun vin ne circulait sans passavant", "Pourquoi")
p.text(s, 0.6, 1.7, 12, 3, ["Au XIXᵉ siècle, le passavant était le laissez-passer des vins.", "", "En 2026, une bouteille vendue en ligne de Vérone à Zurich attend encore… une ressaisie."], 28, INK, True)
p.callout(s, 0.6, 5.2, 12.1, 1.2, "Passavin : le connecteur neutre qui branche la boutique des domaines européens à plusieurs logisticiens vin.", 18, WINE)
# 3
s = p.slide("3 frictions entre vendre et livrer", "Problème")
for j, (t, b) in enumerate([("Pays impossibles", "proposés au panier, au mauvais prix"), ("12 min", "de ressaisie par commande étrangère [hypothèse]"), ("Suivi manuel", "clients inquiets, incidents par mail")]):
    p.box(s, 0.6 + j * 4.1, 1.8, 3.85, 3.0, WHITE); p.text(s, 0.85 + j * 4.1, 2.1, 3.4, 1.0, t, 30, WINE, True); p.text(s, 0.85 + j * 4.1, 3.2, 3.4, 1.4, b, 18)
p.text(s, 0.6, 5.3, 12, 1.2, "Et côté logisticien : chaque intégration refaite à la main, ou un plugin maison peu adopté (< 10 installations).", 18, INK, True)
# 4
s = p.slide("Démonstration : une commande de Vérone vers Zurich", "Solution")
p.table(s, 0.6, 1.6, 12.1, [["Étape", "Aujourd'hui", "Avec Passavin"], ["Panier", "Forfait « Europe » approximatif", "Tarif réel du logisticien, pays non livrables masqués"],
                            ["Après paiement", "Ressaisie dans le portail du logisticien", "Commande transmise automatiquement"], ["Suivi", "Numéro recopié à la main", "Statuts et mail client automatiques"],
                            ["Changer de logisticien", "Tout redévelopper", "Un clic : plusieurs logisticiens branchés"]], widths=[2.6, 4.4, 5.1], size=17, h_row=0.8)
p.text(s, 0.6, 5.9, 12, 0.8, "Règles d'éligibilité fournies et garanties par les logisticiens · pas d'envoi direct aux particuliers américains", 14, GREY)
# 5
s = p.slide("Un marché européen, plus seulement une niche", "Marché")
p.kpis(s, 0.6, 1.6, [(n(TAM_TOT), "TAM : domaines qui vendent en ligne (5 pays)"), (n(SAM_TOT), "SAM : 3 CMS × export aux particuliers"),
                     ("×4 à 7", "par rapport à la V1 (France, WooCommerce)"), (n(Y[2]["clients_fin"]), f"SOM fin 2029 ({Y[2]['clients_fin'] / SAM_TOT:.0%} du SAM)".replace("%", " %"))], w=2.85, big=34)
p.table(s, 0.6, 3.6, 12.1, [["", "France", "Italie", "Espagne", "Allemagne", "Portugal"], ["SAM", *[n(MK[c]["sam_3cms"]) for c in ("FR", "IT", "ES", "DE", "PT")]],
                            ["Domaines M36", *[n(CNT[c]) for c in ("FR", "IT", "ES", "DE", "PT")]], ["Canal", "Sales + réseau", "Sales + réseau", "Sales ES-PT + réseau", "Réseau seul", "Réseau seul"]],
        widths=[2.1, 2.0, 2.0, 2.0, 2.0, 2.0], size=15, h_row=0.6)
p.text(s, 0.6, 6.2, 12, 0.5, "Sources : ICQRF, INE, Destatis, FranceAgriMer, Afnic, WMTips 2026 ; parts d'export = hypothèses à mesurer.", 12, GREY)
# 6
s = p.slide("L'idée existe… et personne n'est neutre", "Concurrence & USP")
p.table(s, 0.6, 1.6, 12.1, [["Acteur", "Limite"], ["Hillebrand Gori eShipping (DHL)", "Mono-logisticien, < 10 installations WooCommerce"], ["MBE Wine for Europe (2026)", "Mono-réseau → client Partner potentiel"],
                            ["VinTrail, Tannico WinePlatform", "Remplacent la boutique du domaine"], ["Sendcloud, Packlink, Boxtal", "Alcool international mal couvert"]], widths=[5.0, 7.1], size=17, h_row=0.7)
p.callout(s, 0.6, 5.3, 12.1, 1.2, "USP : le connecteur NEUTRE qui branche votre boutique à PLUSIEURS logisticiens vin.", 20, WINE)
# 7
s = p.slide("Deux offres, un seul moteur", "Business model")
p.card(s, 0.6, 1.6, 5.9, 3.0, "DIRECT · domaines", ["49 / 99 / 199 €/mois (ARPA 94 €)", "Installation 490 €", "1,50 €/expédition payé par le logisticien (à valider)", f"2029 : {k(Y[2]['ca_direct'])} de CA"], tcolor=WINE, size=16, tsize=20)
p.card(s, 6.8, 1.6, 5.9, 3.0, "PARTNER · logisticiens", ["Licence 690 à 2 490 €/mois", "Mise en place 4 à 12 k€", "0,40 à 0,60 €/expédition", f"2029 : {k(Y[2]['ca_wl'])} de CA, 4 licences"], tcolor=TEAL, size=16, tsize=20)
p.kpis(s, 0.6, 4.9, [(f"{U['ltv']:,.0f} €".replace(",", " "), "LTV d'un domaine"), (f"{U['cac_a3']:,.0f} €".replace(",", " "), "CAC 2029"), (f"{U['ltv_cac']:.0f}", "LTV / CAC"), ("69 k€", "coût d'un plugin maison pour un logisticien (1ʳᵉ année, 2 CMS)")], w=2.85, h=1.5, big=28)
# 8
s = p.slide("Un pays et un CMS à la fois", "Go-to-market Europe")
p.flow(s, 0.6, 1.7, [("2027", f"France · WooCommerce · {n(Y[0]['clients_fin'])} domaines"), ("M12", f"Seed {k(SEED)}"), ("M17-M18", "Italie · Shopify"), ("M20", "Espagne-Portugal"), ("M25", "PrestaShop"), ("2029", f"{n(Y[2]['clients_fin'])} domaines + {n(Y[2]['wl_domaines'])} en marque blanche")], w=1.95, h=1.9, size=14)
p.image(s, CH("pays.png"), 2.4, 3.75, h=3.1)
# 9
s = p.slide(f"Une équipe construite par jalons : {Y[0]['effectif_fin']} → {Y[1]['effectif_fin']} → {Y[2]['effectif_fin']}", "Équipe")
for j, (t, body) in enumerate([("2027", ["Dirigeant·e (commercial transport de vin)", "Lead dev / CTO (M10)", "Sales France (M10)"]),
                               ("2028", ["Partner manager, alternant·e, dév. Shopify", "Customer success FR/EN", "Sales Italie, Sales ES-PT"]),
                               ("2029", ["Customer success IT/ES", "Dév. marque blanche", f"Réseau : {Y[2]['agences_fin']} agences, {Y[2]['apporteurs_fin']} apporteurs, 4 logisticiens"])]):
    p.card(s, 0.6 + j * 4.1, 1.6, 3.85, 3.3, t, body, tcolor=WINE, size=15, tsize=22)
p.callout(s, 0.6, 5.25, 12.1, 1.2, "Mon apport : le terrain et les logisticiens. Ce qui manque : un·e CTO (priorité n°1). Aucun recrutement sans jalon ni 6 mois de trésorerie.", 16)
# 10
s = p.slide("L'équipe commerciale : objectifs et rémunération", "Management — équipe")
p.table(s, 0.6, 1.6, 12.1, [["Poste", "Objectif 2029", "Rémunération"], ["Sales France", f"{n(src(2,'sales_fr'))} domaines", "32 k€ + 15 % de la 1ʳᵉ année signée"],
                            ["Sales Italie", f"{n(src(2,'sales_it'))} domaines", "Idem (Employer of Record)"], ["Sales Espagne & Portugal", f"{n(src(2,'sales_es'))} domaines", "Idem"],
                            ["Partner manager", f"{Y[2]['agences_fin']} agences, 4 licences", "38 k€ + 1 500 €/licence"], ["Customer success ×2", "Résiliation ≤ 1,5 %/mois", "32 k€ dont prime de rétention"]],
        widths=[3.6, 4.0, 4.5], size=17, h_row=0.72)
p.text(s, 0.6, 6.15, 12, 0.6, "Variable récupéré si le client part avant 3 mois · BSPCE · OKR trimestriels", 15, GREY)
# 11
s = p.slide("Le réseau : trois leviers externes", "Management — réseau")
for j, (t, b) in enumerate([("27 agences", ["Silver : 20 % récurrent à vie", "Gold : 25 % + co-marketing", f"{n(AG)} domaines sur 3 ans"]),
                            ("20 apporteurs", ["Consultants, œnologues, syndicats", "250 € par domaine actif à 3 mois", f"{n(tot('apporteur'))} domaines sur 3 ans"]),
                            ("4 logisticiens", ["Revendeurs en marque blanche", "≥ 40 domaines activés/an chacun", f"{n(Y[2]['wl_domaines'])} domaines à M36"])]):
    p.card(s, 0.6 + j * 4.1, 1.6, 3.85, 3.2, t, b, tcolor=TEAL, size=16, tsize=26)
p.callout(s, 0.6, 5.2, 12.1, 1.2, "Recrutés, certifiés, animés (CRM, revues trimestrielles, challenges) — et désactivés sans coût s'ils ne produisent pas.", 16)
# 12
s = p.slide("Manager à distance, dans 5 pays", "Management — leadership")
p.table(s, 0.6, 1.6, 6.3, [["Style situationnel", "Pour qui"], ["Directif", "Alternant·e, sales en rampe"], ["Persuasif", "Sales IT/ES, 6 premiers mois"], ["Participatif", "Customer success, devs"], ["Délégatif", "Sales FR confirmé·e, PM"]], widths=[2.5, 3.8], size=15, h_row=0.65)
p.card(s, 7.2, 1.6, 5.5, 3.3, "Rituels & outils", ["Revue de pipeline hebdo, 1:1 hebdo", "OKR trimestriels reliés au BP", "CRM HubSpot + tableau de bord Metabase", "Challenges : « Pro d'abord », Partenaire de l'année"], size=15, tsize=18)
p.callout(s, 0.6, 5.25, 12.1, 1.2, "Sous-performance : diagnostic → plan de 60 jours → bilan. Des objectifs réalistes et connus à l'avance.", 16)
# 13 — Démo 1
s = p.slide("Démonstration 1 : le BP sort des ressources commerciales", "Business plan")
p.image(s, CH("sources.png"), 0.6, 1.5, w=7.1)
p.bullets(s, 7.9, 1.6, 4.8, 5.0, [f"{n(NEW_TOT)} domaines signés sur 3 ans, chacun attribué à une ressource",
                                  "Productivité × part du vivier couverte par les CMS (47 % → 65 % → 100 %)",
                                  "Rampe de 3-4 mois, saturation du SAM",
                                  "Embauche = jalon + 6 mois de trésorerie",
                                  "**CA = domaines × 94 € + licences : rien sans ressource"], 16, gap=10)
# 14
s = p.slide("Compte de résultat — central", "Business plan")
p.image(s, CH("ca_ebitda.png"), 0.6, 1.5, w=6.6)
p.table(s, 7.5, 1.6, 5.2, [["", "2027", "2028", "2029"], ["CA", k(Y[0]["ca_total"]), k(Y[1]["ca_total"]), k(Y[2]["ca_total"])], ["EBITDA", k(Y[0]["ebitda"]), k(Y[1]["ebitda"]), k(Y[2]["ebitda"])],
                           ["Revenu récurrent annualisé", k(Y[0]["arr_fin"]), k(Y[1]["arr_fin"]), k(Y[2]["arr_fin"])], ["Effectif", n(Y[0]["effectif_fin"]), n(Y[1]["effectif_fin"]), n(Y[2]["effectif_fin"])]],
        widths=[2.0, 1.05, 1.05, 1.1], size=14, h_row=0.6)
p.text(s, 7.5, 4.9, 5.2, 1.5, f"Marge brute {Y[2]['taux_marge_brute']:.0%} · LTV/CAC ≈ {U['ltv_cac']:.0f}".replace("%", " %"), 16, INK, True)
# 15
s = p.slide(f"Point mort : {BE}, pas M36 — et ce qu'il faudrait", "Scénarios")
p.image(s, CH("ebitda_mensuel.png"), 0.6, 1.5, w=6.6)
p.table(s, 7.5, 1.6, 5.2, [["Scénario", "Point mort"], ["Pessimiste", "non atteint → pivot"], ["**Central", f"**{BE}"], ["Optimiste", BE_OPT]], widths=[2.6, 2.6], size=15, h_row=0.55)
p.callout(s, 7.5, 4.05, 5.2, 2.4, "M36 atteint avec : ARPA 99 € (Pro par défaut) + une 5ᵉ licence à M30. Je pilote ces deux leviers.", 15, WINE)
# 16
s = p.slide(f"{k(TOTAL_FUND)} en deux temps, la seed sous conditions", "Besoin financier")
p.kpis(s, 0.6, 1.6, [(k(PRE), "pré-amorçage : apport, prêt d'honneur, BPI, banque"), (k(SEED), f"seed à M{P['seed_month']} : {k(P['seed_equity'])} de fonds propres + {k(P['seed_bpi_loan'])} de prêt BPI"),
                     ("18-22 %", "de dilution (valorisation avant levée de 2-2,5 M€)")], w=3.9, big=34)
p.table(s, 0.6, 3.6, 12.1, [["Conditions GO seed", "≥ 25 domaines payants", "≥ 2 logisticiens intégrés", "≥ 1 lettre d'intention de licence", "Résiliation ≤ 2 %"]], widths=[2.7, 2.3, 2.3, 2.6, 2.2], size=14, h_row=0.9)
p.text(s, 0.6, 4.9, 12, 1.5, ["Affectation : équipe commerciale et réseau 45 % · produit 30 % · marketing Europe 10 % · customer success 5 % · juridique 3 % · réserve 7 %",
                              f"Trésorerie à M36 : {k(Y[2]['tresorerie_fin'])} → pas de nouvelle levée nécessaire pour survivre"], 15)
# 17 — Démo 2
s = p.slide("Démonstration 2 : risques évalués et justifiés", "Risques")
p.table(s, 0.6, 1.5, 12.1, [["Risque", "P×I", "Parade"], ["Les domaines ne paient pas", "12", "8 pré-ventes avant le MVP"], ["Seed non obtenue à M12", "12", "Jalons publics, relais BPI, sinon plan V1"],
                            ["Ouverture IT/ES plus lente", "9", "Recrutement par jalons, rampe de 4 mois"], ["Hypothèses IT/ES non mesurées", "9", "18 entretiens avant le recrutement"],
                            ["Concurrent gratuit (logisticien)", "9", "Neutralité ; en faire un client Partner"], ["Agences peu actives", "9", "Commission à vie, Gold ; sinon 2ᵉ sales FR"]],
        widths=[5.0, 1.2, 5.9], size=16, h_row=0.62)
p.text(s, 0.6, 6.0, 12, 0.8, "17 risques notés · 5 critiques, aucun technique · règle automatique : pas d'embauche sous 6 mois de trésorerie", 15, WINE, True)
# 18 — Démo 3
s = p.slide("Démonstration 3 : le tableau de bord = les objectifs", "Tableaux de bord")
p.table(s, 0.6, 1.5, 12.1, [["Objectif 2029", "KPI", "Responsable", "Si écart"], [f"{n(Y[2]['nouveaux'])} nouveaux domaines", "Signatures par sales (hebdo)", "Chaque sales", "Plan de 60 jours"],
                            [f"{n(Y[2]['clients_fin'])} domaines actifs", "Résiliation ≤ 1,5 %", "Customer success", "Entretien de départ"], [f"{Y[2]['agences_fin']} agences actives", "Domaines via agences", "Partner manager", "Gold, recrutement"],
                            ["4 licences", "Licences signées", "Dirigeant + PM", "Revoir l'offre Partner"], [f"Point mort {BE} → M36", "ARPA (99 €), 5ᵉ licence", "Dirigeant", "Pro par défaut"]],
        widths=[3.2, 3.6, 2.6, 2.7], size=16, h_row=0.68)
p.text(s, 0.6, 6.0, 12, 0.8, "26 KPI · cibles lues dans le même modèle que le BP", 15, GREY)
# 19
s = p.slide("Ce qui reste à prouver — seuils fixés avant les résultats", "Validation")
p.table(s, 0.6, 1.6, 12.1, [["Indicateur", "GO", "STOP"], ["Pré-ventes avec acompte", "≥ 8", "< 4"], ["Logisticiens prêts à signer", "≥ 2 (dont 1 hors Planet Aura)", "0 hors Planet Aura"],
                            ["Prix acceptable médian", "≥ 49 €/mois", "< 25 €"], ["Temps de ressaisie par commande", "≥ 8 min", "< 4 min"], ["Entretiens en Italie", "≥ 3 caves prêtes à tester sur 10", "0"]],
        widths=[5.0, 4.0, 3.1], size=16, h_row=0.7)
p.text(s, 0.6, 6.0, 12, 0.8, "Vos vins passent les frontières. Vos commandes passent toutes seules.", 18, WINE, True)
# 20
s = p.slide("En résumé", "Synthèse et demande")
p.bullets(s, 0.6, 1.6, 12, 3.6, [f"Un vrai irritant, sur un marché européen : {n(SAM_TOT)} domaines adressables dans 5 pays",
                                 "Deux offres sur un seul moteur : Direct (domaines) et Partner (logisticiens)",
                                 f"Une équipe et un réseau qui produisent chaque euro : {Y[2]['effectif_fin']} personnes, {Y[2]['agences_fin']} agences, {Y[2]['apporteurs_fin']} apporteurs, 4 logisticiens",
                                 f"{k(Y[2]['arr_fin'])} de revenu récurrent annualisé en 2029 ; point mort {BE}, M36 avec deux leviers pilotés",
                                 "Des seuils GO / STOP fixés avant les résultats"], 20, gap=12)
p.callout(s, 0.6, 5.3, 12.1, 1.2, f"Notre demande : {k(SEED)} à M{P['seed_month']}, seulement si les jalons de traction sont atteints.", 20, WINE)
assert p.i == 20, p.i
p.save("Passavin_Pitch_Soutenance.pptx")
print("decks ok", d.i, p.i)
